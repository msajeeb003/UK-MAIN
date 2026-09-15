"""
Credit-limits grid (BRD 2.6, S6): buyer rows × one column per insurer with
a quote, over the project's working document (`state.credit`,
`state.columns`, `state.limitsHidden` — the shape the S6 screen stores).

Money rule (shared with the screen): numeric text is normalised to full
pounds "£1,234,567"; zero / nil / declined becomes "0" (a declined limit,
distinct from a blank — BRD 2.5); anything else is kept verbatim so
wording like "Pending" survives. Totals sum the parseable amounts.

Exports: the table as an editable Excel workbook and as a PDF whose value
cells are text form fields (editable in any PDF viewer).
"""

from __future__ import annotations

import io
import re
import secrets
import time
from typing import Any

import openpyxl
import pymupdf
from openpyxl.styles import Alignment, PatternFill
from openpyxl.styles import Font as XlsxFont
from openpyxl.utils import get_column_letter

BUYER_FIELDS = ("buyer", "company_number", "required")
# state.credit rows use the screen's short keys; the API speaks in full names.
_STATE_KEY = {"buyer": "buyer", "company_number": "reg", "required": "req"}

_ZERO_WORDS = re.compile(r"^(?:nil|none|declined|zero|0|£\s*0|0\.00|£\s*0\.00)$", re.I)
_AMOUNT = re.compile(r"^[£$€]?\s*([\d,]+(?:\.\d+)?)\s*$")


class RowNotFound(LookupError):
    pass


class ColumnNotFound(LookupError):
    pass


class NoLimitData(ValueError):
    pass


# ── Money ──────────────────────────────────────────────────────────────────

def parse_amount(value: str) -> int | None:
    """Full pounds, or None when the text is not an amount."""
    v = (value or "").strip()
    if not v:
        return None
    if _ZERO_WORDS.match(v):
        return 0
    m = _AMOUNT.match(v)
    if not m:
        return None
    try:
        return round(float(m.group(1).replace(",", "")))
    except ValueError:
        return None


def format_pounds(n: int) -> str:
    return f"£{n:,}"


def normalise_amount(value: str) -> str:
    n = parse_amount(value)
    if n is None:
        return (value or "").strip()
    return "0" if n == 0 else format_pounds(n)


def is_declined(value: str) -> bool:
    return parse_amount(value) == 0


def total(values: list[str]) -> int | None:
    found, acc = False, 0
    for v in values:
        n = parse_amount(v or "")
        if n is not None:
            acc, found = acc + n, True
    return acc if found else None


# ── Reading ────────────────────────────────────────────────────────────────

def rows(state: dict) -> list[dict]:
    credit = state.get("credit")
    return [r for r in credit if isinstance(r, dict) and r.get("id")] if isinstance(credit, list) else []


def find_row(state: dict, row_id: str) -> dict:
    for r in rows(state):
        if r["id"] == row_id:
            return r
    raise RowNotFound(row_id)


def all_quote_columns(state: dict) -> list[dict]:
    cols = state.get("columns")
    if not isinstance(cols, list):
        return []
    return [c for c in cols if isinstance(c, dict) and c.get("id") and not c.get("expiring")]


def hidden_ids(state: dict) -> list[str]:
    h = state.get("limitsHidden")
    return [x for x in h if isinstance(x, str)] if isinstance(h, list) else []


def columns(state: dict) -> list[dict]:
    """One column per insurer with a quote, minus the ones hidden here."""
    hidden = set(hidden_ids(state))
    return [c for c in all_quote_columns(state) if c["id"] not in hidden]


def find_column(state: dict, column_id: str) -> dict:
    for c in columns(state):
        if c["id"] == column_id:
            return c
    raise ColumnNotFound(column_id)


def row_has_content(r: dict) -> bool:
    return bool((r.get("buyer") or "").strip() or (r.get("reg") or "").strip()
                or (r.get("req") or "").strip()
                or any((v or "").strip() for v in (r.get("offers") or {}).values()))


def provenance(r: dict, key: str, value: str) -> str:
    if not (value or "").strip():
        return "blank"
    if (r.get("edited") or {}).get(key):
        return "edited"
    if r.get("manual"):
        return "manual"
    return "extracted"


def _cell(r: dict, key: str, value: str) -> dict:
    value = value or ""
    return {"value": value, "provenance": provenance(r, key, value),
            "amount": parse_amount(value), "declined": is_declined(value)}


def row_view(r: dict, cols: list[dict]) -> dict:
    offers = r.get("offers") or {}
    return {
        "id": r["id"],
        "manual": bool(r.get("manual")),
        "buyer": _cell(r, "buyer", r.get("buyer") or ""),
        "company_number": _cell(r, "reg", r.get("reg") or ""),
        "required": _cell(r, "req", r.get("req") or ""),
        "offers": {c["id"]: _cell(r, c["id"], offers.get(c["id"]) or "") for c in cols},
        # Offers that arrived before their insurer's column existed (kept
        # by the upload flow; surfaced here so nothing is silently lost).
        "pending": {k: v for k, v in (r.get("pending") or {}).items() if isinstance(v, str)},
    }


def grid_view(project_id: str, state: dict) -> dict:
    cols = columns(state)
    rs = rows(state)
    hidden = set(hidden_ids(state))
    return {
        "project_id": project_id,
        "columns": [{"id": c["id"], "name": str(c.get("name") or ""), "matched": c.get("matched") or None}
                    for c in cols],
        "hidden_columns": [{"id": c["id"], "name": str(c.get("name") or "")}
                           for c in all_quote_columns(state) if c["id"] in hidden],
        "rows": [row_view(r, cols) for r in rs],
        "totals": {
            "required": total([r.get("req") or "" for r in rs]),
            "offers": {c["id"]: total([(r.get("offers") or {}).get(c["id"]) or "" for r in rs]) for c in cols},
        },
        "has_data": any(row_has_content(r) for r in rs),
    }


# ── Mutations (return a NEW state; None = nothing changed) ─────────────────

def _now_ms() -> int:
    return int(time.time() * 1000)


def _with_rows(state: dict, credit: list[dict]) -> dict:
    return {**state, "credit": credit, "updated": _now_ms()}


def _mark_edited(r: dict, key: str) -> dict:
    if r.get("manual"):
        return r
    return {**r, "edited": {**(r.get("edited") or {}), key: True}}


def new_row_id() -> str:
    return secrets.token_hex(6)


def add_row(state: dict, buyer: str = "", company_number: str = "", required: str = "",
            offers: dict[str, str] | None = None) -> tuple[dict, str]:
    """"Add buyer": a broker-entered row (a facility agreed offline)."""
    visible = {c["id"] for c in columns(state)}
    for col_id in offers or {}:
        if col_id not in visible:
            raise ColumnNotFound(col_id)
    row: dict[str, Any] = {
        "id": new_row_id(), "buyer": buyer.strip(), "reg": company_number.strip(),
        "req": normalise_amount(required), "manual": True,
        "offers": {k: normalise_amount(v) for k, v in (offers or {}).items() if normalise_amount(v)},
    }
    return _with_rows(state, [*rows(state), row]), row["id"]


def update_row(state: dict, row_id: str, fields: dict[str, str]) -> dict | None:
    """Change the buyer fields present in `fields` (full names)."""
    r = find_row(state, row_id)
    nxt = r
    for name, value in fields.items():
        key = _STATE_KEY[name]
        clean = normalise_amount(value) if name == "required" else (value or "").strip()
        if (nxt.get(key) or "") == clean:
            continue
        nxt = _mark_edited({**nxt, key: clean}, key)
    if nxt is r:
        return None
    return _with_rows(state, [nxt if x["id"] == row_id else x for x in rows(state)])


def delete_row(state: dict, row_id: str) -> dict:
    find_row(state, row_id)
    return _with_rows(state, [x for x in rows(state) if x["id"] != row_id])


def set_offer(state: dict, row_id: str, column_id: str, value: str) -> dict | None:
    """One insurer's offered limit for a buyer; empty clears it."""
    r = find_row(state, row_id)
    find_column(state, column_id)
    clean = normalise_amount(value)
    offers = dict(r.get("offers") or {})
    if (offers.get(column_id) or "") == clean:
        return None
    if clean:
        offers[column_id] = clean
    else:
        offers.pop(column_id, None)
    nxt = _mark_edited({**r, "offers": offers}, column_id)
    return _with_rows(state, [nxt if x["id"] == row_id else x for x in rows(state)])


def set_column_hidden(state: dict, column_id: str, hidden: bool) -> dict | None:
    """Hide an insurer from the credit-limit table (its offers are dropped
    from every row) or bring it back. The comparison column is untouched."""
    if column_id not in {c["id"] for c in all_quote_columns(state)}:
        raise ColumnNotFound(column_id)
    current = hidden_ids(state)
    if hidden == (column_id in current):
        return None
    if hidden:
        credit = []
        for r in rows(state):
            offers = {k: v for k, v in (r.get("offers") or {}).items() if k != column_id}
            edited = {k: v for k, v in (r.get("edited") or {}).items() if k != column_id}
            nxt = {**r, "offers": offers}
            if edited:
                nxt["edited"] = edited
            else:
                nxt.pop("edited", None)
            credit.append(nxt)
        return {**_with_rows(state, credit), "limitsHidden": [*current, column_id]}
    return {**state, "limitsHidden": [x for x in current if x != column_id], "updated": _now_ms()}


# ── Exports ────────────────────────────────────────────────────────────────

def export_filename(client_name: str, extension: str) -> str:
    client = re.sub(r"[^A-Za-z0-9 \-]", "", client_name or "").strip() or "Client"
    return f"Credit Limits - {client[:80]}.{extension}"


def _table(view: dict) -> tuple[list[str], list[list[str]], list[str]]:
    """Headers, one text row per buyer, and the total row."""
    cols = view["columns"]
    headers = ["Top Customers", "Company Reg", "Limit Required (GBP)"] + [c["name"] for c in cols]
    body = [
        [r["buyer"]["value"], r["company_number"]["value"], r["required"]["value"]]
        + [r["offers"][c["id"]]["value"] for c in cols]
        for r in view["rows"] if row_has_content({
            "buyer": r["buyer"]["value"], "reg": r["company_number"]["value"],
            "req": r["required"]["value"], "offers": {k: v["value"] for k, v in r["offers"].items()},
        })
    ]
    if not body:
        raise NoLimitData("No buyer credit limits to export.")
    t = view["totals"]
    total_row = ["Total", "", format_pounds(t["required"]) if t["required"] is not None else ""]
    total_row += [format_pounds(t["offers"][c["id"]]) if t["offers"][c["id"]] is not None else "" for c in cols]
    return headers, body, total_row


def build_xlsx(view: dict, client_name: str) -> bytes:
    """Editable Excel workbook: values as typed, totals as formulas."""
    headers, body, _ = _table(view)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Credit Limits"
    ws.append([f"Credit limits — {client_name or 'Client'}"])
    ws["A1"].font = XlsxFont(bold=True, size=13)
    ws.append(headers)
    for cell in ws[2]:
        cell.font = XlsxFont(bold=True)
        cell.fill = PatternFill("solid", fgColor="E8ECF4")
    first = 3
    for row in body:
        ws.append(row)
    last = ws.max_row
    ws.append(["Total", ""] + [
        f'=SUMPRODUCT(--SUBSTITUTE(SUBSTITUTE({get_column_letter(i)}{first}:{get_column_letter(i)}{last},"£",""),",",""))'
        if last >= first else ""
        for i in range(3, len(headers) + 1)
    ])
    for cell in ws[ws.max_row]:
        cell.font = XlsxFont(bold=True)
    for i, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(16, len(header) + 4)
        if i >= 3:
            for r in range(first, ws.max_row + 1):
                ws.cell(row=r, column=i).alignment = Alignment(horizontal="right")
    ws.freeze_panes = "A3"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_PAGE_W, _PAGE_H = 842, 595            # A4 landscape, points
_MARGIN, _ROW_H, _HEADER_H = 36, 22, 26
_ROWS_PER_PAGE = 20


def build_pdf(view: dict, client_name: str) -> bytes:
    """Editable PDF: the same table, every value cell a text form field so
    the broker can amend figures in any PDF viewer. Totals are printed
    (a viewer cannot recompute them); the Excel export carries formulas."""
    headers, body, total_row = _table(view)
    n_cols = len(headers)
    usable = _PAGE_W - 2 * _MARGIN
    # Buyer name gets double width; the rest share the remainder.
    unit = usable / (n_cols + 1)
    widths = [unit * 2] + [unit] * (n_cols - 1)
    doc = pymupdf.open()
    field_no = 0

    def header_block(page: pymupdf.Page, title: bool) -> float:
        y = _MARGIN
        if title:
            page.insert_text((_MARGIN, y + 14), f"Credit limits — {client_name or 'Client'}",
                             fontsize=14, fontname="helv")
            page.insert_text((_MARGIN, y + 30), "Editable: click a value to amend it. Amounts in GBP.",
                             fontsize=8.5, fontname="helv", color=(0.4, 0.4, 0.45))
            y += 44
        x = _MARGIN
        for w, h in zip(widths, headers, strict=True):
            rect = pymupdf.Rect(x, y, x + w, y + _HEADER_H)
            page.draw_rect(rect, color=(0.8, 0.82, 0.88), fill=(0.91, 0.93, 0.96), width=0.5)
            page.insert_textbox(rect + (4, 0, -4, 0), h, fontsize=8.5, fontname="hebo", align=0)
            x += w
        return y + _HEADER_H

    for start in range(0, len(body), _ROWS_PER_PAGE):
        page = doc.new_page(width=_PAGE_W, height=_PAGE_H)
        y = header_block(page, title=start == 0)
        chunk = body[start:start + _ROWS_PER_PAGE]
        last_page = start + _ROWS_PER_PAGE >= len(body)
        for row in chunk:
            x = _MARGIN
            for w, value in zip(widths, row, strict=True):
                rect = pymupdf.Rect(x, y, x + w, y + _ROW_H)
                page.draw_rect(rect, color=(0.85, 0.86, 0.9), width=0.5)
                widget = pymupdf.Widget()
                widget.field_name = f"cell_{field_no}"
                widget.field_type = pymupdf.PDF_WIDGET_TYPE_TEXT
                widget.rect = rect + (2, 2, -2, -2)
                # Form-field text is stored in PDFDocEncoding; the pound sign
                # does not round-trip reliably, so amounts are written as
                # plain figures (the header says GBP).
                widget.field_value = value.lstrip("£").strip()
                widget.text_font = "Helv"
                widget.text_fontsize = 8.5
                widget.text_color = (0.1, 0.1, 0.15)
                widget.fill_color = (1, 1, 1)
                widget.border_color = (0.9, 0.9, 0.93)
                page.add_widget(widget)
                field_no += 1
                x += w
            y += _ROW_H
        if last_page:
            x = _MARGIN
            for w, value in zip(widths, total_row, strict=True):
                rect = pymupdf.Rect(x, y, x + w, y + _ROW_H)
                page.draw_rect(rect, color=(0.8, 0.82, 0.88), fill=(0.95, 0.96, 0.98), width=0.5)
                page.insert_textbox(rect + (4, 5, -4, 0), value, fontsize=8.5, fontname="hebo",
                                    align=2 if value.startswith("£") else 0)
                x += w
    doc.set_metadata({"title": f"Credit limits - {client_name or 'Client'}", "producer": "Quote Comparison Tool"})
    out = doc.tobytes(garbage=3, deflate=True)
    doc.close()
    return out

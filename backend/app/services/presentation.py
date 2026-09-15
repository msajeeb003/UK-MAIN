"""Presentation generation (BRD 2.8): editable PPTX + PDF + credit-limit XLSX.

7-page deck matching the brokerage's approved template (cover, About Us,
Feedback of Terms, Terms Comparison, Credit Limits, Demands & needs,
Contact). Server-side rules: 409 until the four key values are confirmed
(2.5); regulatory wording always rendered (2.7); blanks stay blank (2.2);
PPTX uses only plain shapes/tables so it opens cleanly in Google Slides.
"""

import io
import re
from datetime import date

import openpyxl
import pymupdf
from openpyxl.styles import Font as XlsxFont
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from app.core.errors import ExportBlockedError, InvalidDocumentError
from app.models.presentation import PRESENTATION_ROWS, PresentationRequest
from app.models.schemas import CONFIRM_REQUIRED_FIELDS
from app.services.library import presentation_wording
from app.services.money import format_money, render_value, sort_limit_rows

# ── Brand palette (from the approved sample deck) ────────────────────────
NAVY = (0x12, 0x39, 0x5E)
TEAL = (0x14, 0xC0, 0xD5)
INK = (0x11, 0x11, 0x11)
INK2 = (0x44, 0x4A, 0x52)
INK3 = (0x8A, 0x92, 0x9C)
REC_FILL = (0xDF, 0xF6, 0xFA)   # recommended column highlight (teal tint)
EXP_FILL = (0xF0, 0xF0, 0xF0)   # expiring-policy column (grey, per sample)
PANEL = (0xF5, 0xF6, 0xF8)
WHITE = (0xFF, 0xFF, 0xFF)
LINE = (0x30, 0x30, 0x30)

# ── Fixed template wording — configuration, not code ────────────────────
# backend/config/presentation.json (BRD 2.7 / 2.8, rule "Regulatory wording
# is fixed"): re-read on change, never editable out of a generated deck.


def _wording() -> dict:
    return presentation_wording()


def _text(key: str) -> str:
    return str(_wording()[key])


def _important_information() -> str:
    """The static block that closes the Feedback of Terms page (Feedback
    Round 1, D4): verbatim paragraphs, printed under the named declined
    line."""
    return "\n".join(str(p).strip() for p in _wording()["important_information"])


def cover_title(req: PresentationRequest) -> str:
    """BRD: one template, two titles — the front page is the only difference."""
    return (
        "Renewal Credit Insurance Presentation"
        if req.project_type == "renewal"
        else "Credit Insurance Presentation"
    )


def check_export_gate(req: PresentationRequest) -> None:
    """BRD 2.5/2.8: export blocked until the four key values are confirmed.
    A limits-only project has no quote columns and nothing to confirm (D1)."""
    if not req.columns:
        return
    missing = [f for f in CONFIRM_REQUIRED_FIELDS if f not in req.confirmed_fields]
    if missing:
        raise ExportBlockedError(
            "Export is blocked until estimated annual premium, indemnity, "
            "excess and max annual liability are confirmed on the review "
            f"screen ({len(missing)} remaining)."
        )


def declined_insurers(req: PresentationRequest) -> list[str]:
    """
    BRD 2.1/S8: approached insurers with no quote column are auto-named as
    declined; a late quote (new column) moves them off this line.

    A column counts as that insurer's quote when the approached name
    matches its standing-list `matched` name OR its display name — a
    column titled with the document's own wording ("HCC International
    Insurance Company plc") must never leave Tokio Marine HCC on the
    declined line.
    """
    matched_names = {c.matched.lower() for c in req.columns if c.matched}
    column_names = [c.name.lower() for c in req.columns]
    declined = []
    for name in req.approached_insurers:
        needle = name.lower()
        quoted = needle in matched_names or any(
            needle in cn or cn in needle for cn in column_names
        )
        if not quoted:
            declined.append(name)
    return declined


# ── Shared content assembly — one source of truth for both renderers ────

def _cover_heading(req: PresentationRequest) -> str:
    return f"{req.client_name} - {cover_title(req)}"


def _subtitle(req: PresentationRequest) -> str:
    # BRD 2.8 cover: client name, date and the title — nothing else.
    return date.today().strftime("%B %Y")


def _approached_sentence(req: PresentationRequest) -> str:
    names = req.approached_insurers or [c.name for c in req.columns]
    if not names:
        return ""
    return (
        "We approached the following insurers to quote for your business: "
        + ", ".join(names) + "."
    )


def _declined_sentence(req: PresentationRequest) -> str | None:
    declined = declined_insurers(req)
    if not declined:
        return None
    verb = "was" if len(declined) == 1 else "were"
    return (
        f"{', '.join(declined)} {verb} approached but declined to quote "
        "— further details are available if requested."
    )


def _recommended_name(req: PresentationRequest) -> str:
    col = next((c for c in req.columns if _is_rec(req, c)), None)
    return col.name if col else "[no insurer selected]"


def _has_recommendation(req: PresentationRequest) -> bool:
    """S7: the broker may continue without recommending. The fixed FCA
    wording (demands & needs, our status) always renders; only the
    "Our recommendation" paragraph and the reasons are left out, so an
    unmerged placeholder never reaches a client deck."""
    return any(_is_rec(req, c) for c in req.columns)


def _is_rec(req: PresentationRequest, col) -> bool:
    return col.id == req.recommended_id


def _is_expiring(col) -> bool:
    return col.name.lower().startswith("expiring")


def _col_fill(req: PresentationRequest, col):
    if _is_rec(req, col):
        return REC_FILL
    if _is_expiring(col):
        return EXP_FILL
    return None


def _reason_lines(req: PresentationRequest) -> list[str]:
    # Strip only a leading enumeration ("1. ", "2) ", "- ") — never the
    # start of the reason itself ("50% cheaper" must stay intact).
    lines = [re.sub(r"^\s*(?:\d+[.)]\s+|[-•]\s+)", "", ln).strip()
             for ln in req.reasons.splitlines() if ln.strip()]
    return [f"{i}. {ln}" for i, ln in enumerate(lines, start=1)]


def _limit_columns(req: PresentationRequest) -> list:
    """Insurer columns of the credit-limit page: the quotes only — the
    expiring policy has no column or row there (Feedback Round 1, D3)."""
    return [c for c in req.columns if not _is_expiring(c)]


def _limit_rows(req: PresentationRequest) -> list:
    """Buyer rows, largest limit required first (D2)."""
    return sort_limit_rows(list(req.credit_limits), lambda r: r.required)


def _limits_headers(req: PresentationRequest) -> list[str]:
    return (["Top Customers", "Company Reg", "Limit Required (GBP)"]
            + [c.name for c in _limit_columns(req)])


def _money_total(values: list[str]) -> str:
    """Sum parseable amounts for the sample deck's Total row; blank if none."""
    total, found = 0, False
    for v in values:
        digits = re.sub(r"[^\d]", "", v or "")
        if digits:
            total += int(digits)
            found = True
    return f"£{total:,}" if found else ""


def _limits_total_row(req: PresentationRequest) -> list[str]:
    rows = _limit_rows(req)
    return (["Total", "", _money_total([r.required for r in rows])]
            + [_money_total([r.offers.get(c.id, "") for r in rows])
               for c in _limit_columns(req)])


def suggested_filename(req: PresentationRequest, extension: str) -> str:
    """S8 download name: "{Renewal|Credit Insurance} Presentation of Terms -
    {Client}.{ext}". ASCII-only: HTTP headers are latin-1 and a non-ASCII
    client name must never be able to break the download response."""
    client = re.sub(r"[^A-Za-z0-9 \-]", "", req.client_name).strip() or "Client"
    kind = "Renewal" if req.project_type == "renewal" else "Credit Insurance"
    return f"{kind} Presentation of Terms - {client[:80]}.{extension}"


def _cell_text(value: str | None) -> str:
    # BRD 2.2: blank stays blank — never a placeholder.
    return (value or "").strip()


# ═════════════════════════════ PPTX ═════════════════════════════════════

def _rgb(t: tuple[int, int, int]) -> RGBColor:
    return RGBColor(*t)


def _add_textbox(slide, x, y, w, h, text, *, size, color=INK, bold=False,
                 align=None, wrap=True):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = wrap
    paragraph = frame.paragraphs[0]
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = _rgb(color)
    return box


def _add_accent(slide, x, y, w, h, color, rotation=18):
    """A rotated brand parallelogram — the sample deck's diagonal motif."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.PARALLELOGRAM, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.rotation = rotation
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(color)
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _brand_mark(slide, x=0.7, y=0.45):
    _add_textbox(slide, x, y, 5.5, 0.8, _wording()["brand"]["name"], size=17, color=NAVY, bold=True)


def _style_cell(cell, text, *, size, bold=False, color=INK, fill=None):
    cell.text_frame.word_wrap = True
    paragraph = cell.text_frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = _rgb(color)
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = _rgb(fill)
    else:
        cell.fill.solid()
        cell.fill.fore_color.rgb = _rgb(WHITE)


def build_pptx(req: PresentationRequest) -> bytes:
    """Render the full deck as an editable, Google-Slides-clean .pptx."""
    check_export_gate(req)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # ── 1. Cover ─────────────────────────────────────────────────────────
    slide = prs.slides.add_slide(blank)
    _add_accent(slide, 9.2, -1.2, 2.2, 10.0, TEAL)
    _add_accent(slide, 10.6, -1.5, 3.4, 10.5, NAVY)
    _brand_mark(slide, 0.7, 0.9)
    _add_textbox(slide, 0.7, 2.6, 8.2, 2.3, _cover_heading(req),
                 size=36, color=INK, bold=True)
    _add_textbox(slide, 0.7, 5.0, 8.0, 0.5, _subtitle(req), size=14, color=INK2)
    _add_textbox(slide, 0.7, 6.35, 8.0, 0.4, "Visit Our Website", size=11, color=INK2)
    _add_textbox(slide, 0.7, 6.7, 8.0, 0.5, _wording()["brand"]["site"], size=16, color=INK, bold=True)

    # ── 2. About Us ──────────────────────────────────────────────────────
    slide = prs.slides.add_slide(blank)
    _add_accent(slide, -1.6, -1.4, 3.6, 10.5, NAVY, rotation=15)
    _add_accent(slide, 2.1, -1.0, 1.6, 9.5, TEAL, rotation=15)
    _add_textbox(slide, 6.2, 0.8, 6.3, 0.7, "About Us", size=30, bold=True)
    _add_textbox(slide, 6.2, 1.7, 6.4, 5.2, _text("about"), size=14, color=INK2)

    # ── 3. Feedback of Terms ─────────────────────────────────────────────
    slide = prs.slides.add_slide(blank)
    _add_textbox(slide, 0.8, 0.35, 11.7, 0.55, "Feedback of Terms", size=24, bold=True)
    _add_textbox(slide, 0.8, 1.0, 11.7, 0.7, _text("feedback_intro"), size=10.5, color=INK2)
    _add_textbox(slide, 0.8, 1.75, 11.7, 1.7, _text("fair_presentation"),
                 size=9.5, color=INK2)
    _add_textbox(slide, 0.8, 3.5, 11.7, 0.3, "TERMS", size=10.5, color=INK3, bold=True)
    _add_textbox(slide, 0.8, 3.8, 11.7, 0.5, _text("terms_notes"), size=9.5, color=INK2)
    _add_textbox(slide, 0.8, 4.35, 11.7, 0.4, _approached_sentence(req), size=10.5)
    declined_line = _declined_sentence(req)
    y = 4.8
    if declined_line:
        _add_textbox(slide, 0.8, y, 11.7, 0.4, declined_line,
                     size=10.5, color=NAVY, bold=True)
        y += 0.45
    _add_textbox(slide, 0.8, y, 11.7, 2.0, _important_information(), size=9.5, color=INK2)

    # ── 4. Terms Comparison ──────────────────────────────────────────────
    slide = prs.slides.add_slide(blank)
    _add_textbox(slide, 0.6, 0.25, 12.1, 0.55, "Terms Comparison", size=24,
                 bold=True, align=None)
    n_cols = len(req.columns)
    n_rows = 1 + len(PRESENTATION_ROWS)
    table_w = 12.1
    label_w = 3.0
    value_w = (table_w - label_w) / max(n_cols, 1)
    shape = slide.shapes.add_table(
        n_rows, 1 + n_cols, Inches(0.6), Inches(0.9), Inches(table_w), Inches(5.8)
    )
    table = shape.table
    table.columns[0].width = Inches(label_w)
    for i in range(n_cols):
        table.columns[i + 1].width = Inches(value_w)

    _style_cell(table.cell(0, 0), "Insurer", size=10.5, bold=True, fill=PANEL)
    for c, col in enumerate(req.columns, start=1):
        _style_cell(table.cell(0, c), col.name, size=10.5, bold=True,
                    color=NAVY if _is_rec(req, col) else INK,
                    fill=_col_fill(req, col) or PANEL)
    for r, (key, label) in enumerate(PRESENTATION_ROWS, start=1):
        _style_cell(table.cell(r, 0), label, size=9, bold=True)
        for c, col in enumerate(req.columns, start=1):
            _style_cell(table.cell(r, c), render_value(key, col.values.get(key)),
                        size=9, fill=_col_fill(req, col))
    if req.notes.strip():
        _add_textbox(slide, 0.6, 6.85, 12.1, 0.55, req.notes.strip(),
                     size=9, color=INK2)

    # ── 5. Credit Limits (omitted cleanly when none) ─────────────────────
    if req.credit_limits:
        limit_cols = _limit_columns(req)
        limit_rows = _limit_rows(req)
        slide = prs.slides.add_slide(blank)
        _add_textbox(slide, 0.6, 0.35, 12, 0.6, "Credit Limits", size=24, bold=True)
        rows = 2 + len(limit_rows)  # header + buyers + Total
        cols = 3 + len(limit_cols)
        shape = slide.shapes.add_table(
            rows, cols, Inches(0.6), Inches(1.1), Inches(12.1),
            Inches(min(5.9, 0.34 * rows)),
        )
        table = shape.table
        for c, header in enumerate(_limits_headers(req)):
            col_obj = limit_cols[c - 3] if c >= 3 else None
            fill = _col_fill(req, col_obj) if col_obj is not None else None
            _style_cell(table.cell(0, c), header, size=10, bold=True,
                        color=NAVY if col_obj is not None and _is_rec(req, col_obj) else INK,
                        fill=fill or PANEL)
        for r, row in enumerate(limit_rows, start=1):
            _style_cell(table.cell(r, 0), _cell_text(row.buyer), size=9.5)
            _style_cell(table.cell(r, 1), _cell_text(row.company_number), size=9.5)
            _style_cell(table.cell(r, 2), format_money(row.required), size=9.5)
            for c, col in enumerate(limit_cols, start=3):
                _style_cell(table.cell(r, c), format_money(row.offers.get(col.id)),
                            size=9.5, fill=_col_fill(req, col))
        for c, text in enumerate(_limits_total_row(req)):
            _style_cell(table.cell(rows - 1, c), text, size=9.5, bold=True,
                        fill=PANEL)

    # ── 6. Demands & needs / recommendation ──────────────────────────────
    slide = prs.slides.add_slide(blank)
    _add_textbox(slide, 0.8, 0.4, 11.7, 0.6,
                 "Our understanding of your demands & needs:", size=22, bold=True)
    _add_textbox(slide, 0.8, 1.1, 11.7, 1.3, _text("demands"), size=11.5, color=INK2)
    if _has_recommendation(req):
        _add_textbox(slide, 0.8, 2.55, 11.7, 0.4, "Our recommendation:", size=14, bold=True)
        _add_textbox(slide, 0.8, 3.0, 11.7, 1.5,
                     _text("recommendation").format(name=_recommended_name(req)),
                     size=11.5, color=INK2)
    reasons = _reason_lines(req) if _has_recommendation(req) else []
    if reasons:
        _add_textbox(slide, 0.8, 4.55, 11.7, 0.4,
                     "In addition to this, the following points were important "
                     "with our recommendation:", size=11.5, bold=True)
        _add_textbox(slide, 0.8, 5.0, 11.7, 1.0, "\n".join(reasons), size=11.5)
    _add_textbox(slide, 0.8, 6.05, 11.7, 0.35, "Our status:", size=12, bold=True)
    _add_textbox(slide, 0.8, 6.4, 11.7, 1.0, _text("our_status"), size=9, color=INK2)

    # ── 7. Contact Us ────────────────────────────────────────────────────
    slide = prs.slides.add_slide(blank)
    _add_accent(slide, 10.4, -1.4, 3.6, 10.5, TEAL, rotation=-15)
    _add_accent(slide, 9.3, -1.0, 1.4, 9.8, NAVY, rotation=-15)
    _add_textbox(slide, 0.8, 0.8, 7.5, 0.7, "Contact Us", size=30, bold=True)
    _add_textbox(slide, 0.8, 1.9, 7.6, 1.4, _text("contact"), size=13.5, color=INK2)
    for i, line in enumerate(_wording()["contact_lines"]):
        pill = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.6 + i * 0.85),
            Inches(4.6), Inches(0.6),
        )
        pill.fill.solid()
        pill.fill.fore_color.rgb = _rgb(NAVY)
        pill.line.fill.background()
        pill.shadow.inherit = False
        frame = pill.text_frame
        frame.word_wrap = False
        run = frame.paragraphs[0].add_run()
        run.text = line
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = _rgb(WHITE)

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


# ═════════════════════════════ PDF ══════════════════════════════════════

PAGE_W, PAGE_H = 960, 540  # 16:9 points
MARGIN = 54

_PDF_CHAR_MAP = str.maketrans({
    "—": "-", "–": "-",
    "‘": "'", "’": "'",
    "“": '"', "”": '"',
    "…": "...", "·": "-", "•": "-",
})


def _norm(color: tuple[int, int, int]) -> tuple[float, float, float]:
    return tuple(v / 255 for v in color)


def _pdf_safe(text: str) -> str:
    """Base-14 Helvetica is Latin-1 only — map common typographic chars."""
    return text.translate(_PDF_CHAR_MAP)


def _pdf_text(page, x, y, w, h, text, *, size, color=INK, bold=False):
    # insert_textbox silently drops ALL text when it doesn't fit its rect.
    # Shrink the font first; as a last resort extend the rect by the
    # reported deficit — overflowing text beats vanished text (BRD 2.2:
    # a broker-written note must reach the export).
    h = max(h, size * 2)
    font = "hebo" if bold else "helv"
    safe = _pdf_safe(text)
    fs = size
    rc = -1.0
    while fs >= 6:
        rc = page.insert_textbox(
            pymupdf.Rect(x, y, x + w, y + h), safe,
            fontsize=fs, fontname=font, color=_norm(color),
        )
        if rc >= 0:
            return
        fs -= 0.5
    page.insert_textbox(
        pymupdf.Rect(x, y, x + w, y + h - rc + 12), safe,
        fontsize=6, fontname=font, color=_norm(color),
    )


def _pdf_accent(page, points, color):
    page.draw_polyline(
        [pymupdf.Point(*p) for p in points],
        color=None, fill=_norm(color), closePath=True,
    )


def _pdf_cover_accents(page, right=True):
    """The sample deck's diagonal navy/teal motif along one edge."""
    if right:
        _pdf_accent(page, [(720, 0), (820, 0), (620, PAGE_H), (520, PAGE_H)], TEAL)
        _pdf_accent(page, [(840, 0), (PAGE_W, 0), (PAGE_W, PAGE_H), (640, PAGE_H)], NAVY)
    else:
        _pdf_accent(page, [(140, 0), (240, 0), (440, PAGE_H), (340, PAGE_H)], TEAL)
        _pdf_accent(page, [(0, 0), (120, 0), (320, PAGE_H), (0, PAGE_H)], NAVY)


def _cell_lines(text: str, width: float, font_size: float, bold: bool) -> int:
    """Rough line count for a wrapped cell — used to size rows so
    insert_textbox can never silently drop long values (BRD 2.2: every
    reviewed value must reach the export)."""
    text = _pdf_safe(text or "")
    if not text:
        return 1
    avail = max(width - 10, 12) * 0.95  # padding + word-wrap slack
    font = "hebo" if bold else "helv"
    lines = 0
    for para in text.split("\n"):
        w = pymupdf.get_text_length(para, fontname=font, fontsize=font_size)
        lines += max(1, -(-int(w) // int(avail)))
    return lines


def _pdf_table(page, x, y, widths, cells, *, font_size, min_row_h=22,
               max_y=PAGE_H - 24):
    """cells: rows of (text, fill|None, bold, color) tuples.

    Row heights grow to fit wrapped text; if the table would overflow the
    page, the font steps down until it fits. Returns the y below the table.
    """
    fs = font_size
    while True:
        heights = []
        for row in cells:
            h = float(min_row_h)
            for c, (text, _fill, bold, _color) in enumerate(row):
                n = _cell_lines(text, widths[c], fs, bold)
                h = max(h, n * fs * 1.35 + 8)
            heights.append(h)
        if y + sum(heights) <= max_y or fs <= 6:
            break
        fs -= 0.5

    cy = y
    for r, row in enumerate(cells):
        cx = x
        for c, (text, fill, bold, color) in enumerate(row):
            rect = pymupdf.Rect(cx, cy, cx + widths[c], cy + heights[r])
            if fill is not None:
                page.draw_rect(rect, color=None, fill=_norm(fill))
            page.draw_rect(rect, color=_norm(LINE), width=0.5)
            # Shrink-to-fit last resort: insert_textbox writes nothing and
            # returns a negative deficit when the text still doesn't fit.
            tfs = fs
            while tfs >= 5:
                rc = page.insert_textbox(
                    rect + (4, 3, -4, -1), _pdf_safe(text), fontsize=tfs,
                    fontname="hebo" if bold else "helv", color=_norm(color),
                )
                if rc >= 0:
                    break
                tfs -= 0.5
            cx += widths[c]
        cy += heights[r]
    return cy


def build_pdf(req: PresentationRequest) -> bytes:
    """Render the same page sequence as the PPTX, as a PDF."""
    check_export_gate(req)
    doc = pymupdf.open()

    # ── 1. Cover ─────────────────────────────────────────────────────────
    page = doc.new_page(width=PAGE_W, height=PAGE_H)
    _pdf_cover_accents(page, right=True)
    _pdf_text(page, MARGIN, 60, 420, 26, _wording()["brand"]["name"], size=15, color=NAVY, bold=True)
    _pdf_text(page, MARGIN, 190, 560, 110, _cover_heading(req), size=27, bold=True)
    _pdf_text(page, MARGIN, 330, 500, 24, _subtitle(req), size=12, color=INK2)
    _pdf_text(page, MARGIN, 430, 400, 18, "Visit Our Website", size=10, color=INK2)
    _pdf_text(page, MARGIN, 450, 400, 22, _wording()["brand"]["site"], size=13, bold=True)

    # ── 2. About Us ──────────────────────────────────────────────────────
    page = doc.new_page(width=PAGE_W, height=PAGE_H)
    _pdf_cover_accents(page, right=False)
    _pdf_text(page, 480, 60, 420, 34, "About Us", size=22, bold=True)
    _pdf_text(page, 480, 110, 420, 360, _text("about"), size=11.5, color=INK2)

    # ── 3. Feedback of Terms ─────────────────────────────────────────────
    page = doc.new_page(width=PAGE_W, height=PAGE_H)
    _pdf_text(page, MARGIN, 30, 840, 28, "Feedback of Terms", size=18, bold=True)
    _pdf_text(page, MARGIN, 66, 852, 38, _text("feedback_intro"), size=9, color=INK2)
    _pdf_text(page, MARGIN, 108, 852, 150, _text("fair_presentation"), size=8.5, color=INK2)
    _pdf_text(page, MARGIN, 262, 840, 14, "TERMS", size=9, color=INK3, bold=True)
    _pdf_text(page, MARGIN, 278, 852, 30, _text("terms_notes"), size=8.5, color=INK2)
    _pdf_text(page, MARGIN, 322, 852, 20, _approached_sentence(req), size=9.5)
    declined_line = _declined_sentence(req)
    y = 350
    if declined_line:
        _pdf_text(page, MARGIN, y, 852, 20, declined_line,
                  size=9.5, color=NAVY, bold=True)
        y += 26
    _pdf_text(page, MARGIN, y, 852, 130, _important_information(), size=8.5, color=INK2)

    # ── 4. Terms Comparison ──────────────────────────────────────────────
    page = doc.new_page(width=PAGE_W, height=PAGE_H)
    _pdf_text(page, MARGIN, 22, 852, 30, "Terms Comparison", size=17, bold=True)
    n_cols = len(req.columns)
    label_w = 200.0
    value_w = (PAGE_W - 2 * MARGIN - label_w) / max(n_cols, 1)
    widths = [label_w] + [value_w] * n_cols
    header = [("Insurer", PANEL, True, INK)] + [
        (c.name, _col_fill(req, c) or PANEL, True,
         NAVY if _is_rec(req, c) else INK)
        for c in req.columns
    ]
    body = []
    for key, label in PRESENTATION_ROWS:
        row = [(label, None, True, INK)]
        for col in req.columns:
            row.append((render_value(key, col.values.get(key)), _col_fill(req, col),
                        False, INK))
        body.append(row)
    end_y = _pdf_table(page, MARGIN, 56, widths, [header] + body,
                       font_size=8, max_y=PAGE_H - 40)
    if req.notes.strip():
        _pdf_text(page, MARGIN, end_y + 8, 852, 30, req.notes.strip(),
                  size=8, color=INK2)

    # ── 5. Credit Limits (omitted cleanly when none) ─────────────────────
    if req.credit_limits:
        limit_cols = _limit_columns(req)
        limit_rows = _limit_rows(req)
        page = doc.new_page(width=PAGE_W, height=PAGE_H)
        _pdf_text(page, MARGIN, 30, 852, 30, "Credit Limits", size=17, bold=True)
        fixed = [200.0, 100.0, 120.0]
        value_w = (PAGE_W - 2 * MARGIN - sum(fixed)) / max(len(limit_cols), 1)
        widths = fixed + [value_w] * len(limit_cols)
        header = [(h, PANEL, True,
                   NAVY if i >= 3 and _is_rec(req, limit_cols[i - 3]) else INK)
                  for i, h in enumerate(_limits_headers(req))]
        for i, col in enumerate(limit_cols, start=3):
            fill = _col_fill(req, col)
            if fill:
                header[i] = (header[i][0], fill, True, header[i][3])
        body = []
        for row in limit_rows:
            cells = [(_cell_text(row.buyer), None, False, INK),
                     (_cell_text(row.company_number), None, False, INK2),
                     (format_money(row.required), None, False, INK)]
            for col in limit_cols:
                cells.append((format_money(row.offers.get(col.id)),
                              _col_fill(req, col), False, INK))
            body.append(cells)
        total = [(t, PANEL, True, INK) for t in _limits_total_row(req)]
        _pdf_table(page, MARGIN, 62, widths, [header] + body + [total],
                   font_size=8.5, min_row_h=20)

    # ── 6. Demands & needs / recommendation ──────────────────────────────
    page = doc.new_page(width=PAGE_W, height=PAGE_H)
    _pdf_text(page, MARGIN, 30, 852, 30,
              "Our understanding of your demands & needs:", size=16, bold=True)
    _pdf_text(page, MARGIN, 70, 852, 74, _text("demands"), size=10, color=INK2)
    if _has_recommendation(req):
        _pdf_text(page, MARGIN, 160, 852, 18, "Our recommendation:", size=11, bold=True)
        _pdf_text(page, MARGIN, 182, 852, 90,
                  _text("recommendation").format(name=_recommended_name(req)),
                  size=10, color=INK2)
    reasons = _reason_lines(req) if _has_recommendation(req) else []
    if reasons:
        _pdf_text(page, MARGIN, 286, 852, 18,
                  "In addition to this, the following points were important "
                  "with our recommendation:", size=10, bold=True)
        _pdf_text(page, MARGIN, 308, 852, 70, "\n".join(reasons), size=10)
    _pdf_text(page, MARGIN, 396, 852, 16, "Our status:", size=10, bold=True)
    _pdf_text(page, MARGIN, 414, 852, 90, _text("our_status"), size=8.5, color=INK2)

    # ── 7. Contact Us ────────────────────────────────────────────────────
    page = doc.new_page(width=PAGE_W, height=PAGE_H)
    _pdf_cover_accents(page, right=True)
    _pdf_text(page, MARGIN, 60, 500, 34, "Contact Us", size=22, bold=True)
    _pdf_text(page, MARGIN, 120, 480, 70, _text("contact"), size=11.5, color=INK2)
    for i, line in enumerate(_wording()["contact_lines"]):
        y = 240 + i * 58
        page.draw_rect(pymupdf.Rect(MARGIN, y, MARGIN + 300, y + 36),
                       color=None, fill=_norm(NAVY))
        _pdf_text(page, MARGIN + 16, y + 10, 280, 20, line, size=12,
                  color=WHITE, bold=True)

    data = doc.tobytes()
    doc.close()
    return data


# ═════════════════════════════ XLSX ═════════════════════════════════════

def build_limits_xlsx(req: PresentationRequest) -> bytes:
    """BRD 2.6 output: the buyer credit-limit table as an editable Excel file."""
    check_export_gate(req)
    if not req.credit_limits:
        raise InvalidDocumentError(
            "No buyer credit limits to export — the credit-limit page is "
            "omitted when no limits are supplied."
        )
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Credit Limits"
    headers = _limits_headers(req)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = XlsxFont(bold=True)
    limit_cols = _limit_columns(req)
    for row in _limit_rows(req):
        sheet.append(
            [row.buyer, row.company_number, format_money(row.required)]
            + [format_money(row.offers.get(c.id, "")) for c in limit_cols]
        )
    sheet.append(_limits_total_row(req))
    for cell in sheet[sheet.max_row]:
        cell.font = XlsxFont(bold=True)
    for i, header in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = max(14, len(header) + 4)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()

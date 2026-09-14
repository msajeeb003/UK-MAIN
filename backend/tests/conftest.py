"""Shared fixtures for the test suite. No network calls, no API keys —
the OpenAI step is monkeypatched so tests exercise everything around it."""

import io
import os

import openpyxl
import pymupdf
import pytest
from fastapi.testclient import TestClient

from app.models.schemas import BuyerCreditLimit, QuoteExtraction, SourcedValue

TEST_USER = ("broker@test.local", "correct-horse-9")


@pytest.fixture(scope="session", autouse=True)
def _isolated_data_dir(tmp_path_factory):
    """Point the SQLite database and file storage at a per-run temp dir."""
    os.environ["DATA_DIR"] = str(tmp_path_factory.mktemp("data"))
    from app.core.config import get_settings
    get_settings.cache_clear()
    yield


@pytest.fixture
def anon_client() -> TestClient:
    """A client with no session — for testing the auth guard itself."""
    from app.main import app
    return TestClient(app)


def sign_in(test_client: TestClient) -> TestClient:
    """Ensure the shared test user exists and start a session."""
    from app.core import auth, db
    if not db.query_one("SELECT id FROM users WHERE email=?", (TEST_USER[0],)):
        auth.create_user(*TEST_USER)
    res = test_client.post(
        "/auth/login", json={"email": TEST_USER[0], "password": TEST_USER[1]}
    )
    assert res.status_code == 200
    # Send the session's CSRF token on every subsequent write, as the real
    # frontend does (X-CSRF-Token header).
    test_client.headers["X-CSRF-Token"] = res.json()["csrf_token"]
    return test_client


@pytest.fixture
def client() -> TestClient:
    """A signed-in client (endpoints require a session — BRD 2.10)."""
    from app.main import app
    return sign_in(TestClient(app))


def make_pdf(pages: list[list[str]]) -> bytes:
    """Build a digital PDF: one inner list of text lines per page."""
    doc = pymupdf.open()
    for lines in pages:
        page = doc.new_page()
        for i, line in enumerate(lines):
            page.insert_text((72, 100 + i * 22), line)
    data = doc.tobytes()
    doc.close()
    return data


def make_xlsx(sheets: dict[str, list[list]]) -> bytes:
    """Build an .xlsx workbook: sheet title -> rows of cell values."""
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets.items():
        sheet = workbook.create_sheet(title)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_xls(sheets: dict[str, list[list]]) -> bytes:
    """Build a LEGACY .xls (OLE2) workbook — real insurers still send these."""
    import xlwt

    workbook = xlwt.Workbook()
    for title, rows in sheets.items():
        sheet = workbook.add_sheet(title)
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                sheet.write(r, c, value)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


DIGITAL_PAGE = [
    "ACME Credit Insurance plc — Quotation",
    "Insurable Turnover: GBP 12,000,000 for the forthcoming policy period.",
    "Premium Rate: 0.055% of insured turnover, payable quarterly in advance.",
    "Indemnity: 90% of each insured loss following the policy excess.",
    "This quotation is valid for 30 days from the date of issue and is",
    "subject to satisfactory completion of a proposal form and to the",
    "insurer's standard policy terms, conditions and exclusions.",
]


@pytest.fixture
def digital_pdf() -> bytes:
    """Two text-rich pages — classifies as DIGITAL."""
    return make_pdf([DIGITAL_PAGE, DIGITAL_PAGE])


@pytest.fixture
def sparse_pdf() -> bytes:
    """Two nearly-empty pages — classifies as SCANNED."""
    return make_pdf([["Page 1"], ["Page 2"]])


@pytest.fixture
def limits_xlsx() -> bytes:
    """A small buyer credit-limit schedule as Excel."""
    return make_xlsx({
        "Limits": [
            ["Buyer Name", "Company registration number", "Application Amount", "Amount Agreed"],
            ["Example Ltd", "01234567", 250000, 200000],
            ["Sample Trading", "07654321", 80000, 80000],
        ],
    })


def sv(value, page, confidence="high") -> SourcedValue:
    if value is None:
        confidence = None
    return SourcedValue(value=value, page=page, confidence=confidence)


@pytest.fixture
def sample_extraction() -> QuoteExtraction:
    """A plausible LLM result, including deliberately bad page citations,
    an uncertain value, and a missing confidence flag."""
    return QuoteExtraction(
        document_type="insurer_quote",
        insurer=sv("ACME Credit Insurance plc", 1),
        annual_turnover=sv("GBP 12,000,000", 1),
        premium_rate=sv("0.055%", 1),
        estimated_annual_premium_exc_ipt=sv(None, 3),   # null value w/ page
        minimum_annual_premium=sv(None, None),
        credit_limit_charges=sv(None, None),
        indemnity=sv("90%", 99),                        # out-of-range page
        excess=SourcedValue(value="GBP 5,000", page=2, confidence=None),  # unflagged
        excess_type=sv("Minimum Retention Each and Every Loss", 2),
        max_annual_liability=sv(None, None),
        discretionary_limit=sv("GBP 20,000", 2, confidence="uncertain"),
        max_terms_of_payment=sv(None, None),
        max_extension_period=sv(None, None),
        additional_info=sv(None, None),
        buyer_credit_limits=[BuyerCreditLimit(
            buyer_name="Example Ltd", company_number="01234567",
            limit_required="GBP 250,000", limit_offered="GBP 200,000", page=42,
        )],
    )

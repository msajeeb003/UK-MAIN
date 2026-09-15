"""The generated deck reproduces the client's approved template: page size,
artwork pictures, fonts, table geometry and borders (BRD 2.8 "Template:
follows the existing presentation design"; Google Slides compatible)."""

import io

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt

from app.services.presentation import ASSETS, REC_FILL, build_pptx
from tests.test_presentation import make_request


def _deck(**overrides):
    return Presentation(io.BytesIO(build_pptx(make_request(**overrides))))


def _pictures(slide):
    return [s for s in slide.shapes if s.shape_type is not None and s.shape_type == 13]  # PICTURE


def _runs(slide):
    return [r for s in slide.shapes if s.has_text_frame for p in s.text_frame.paragraphs for r in p.runs]


def test_template_assets_are_shipped():
    for name in ("bg-cover.png", "bg-about.png", "bg-demands.png", "bg-contact.png", "corner.png"):
        assert (ASSETS / name).stat().st_size > 10_000, name
    fonts = ASSETS.parent / "fonts"
    for name in ("Poppins-Regular.ttf", "Poppins-Bold.ttf", "Antonio-Bold.ttf"):
        assert (fonts / name).read_bytes()[:4] == b"\x00\x01\x00\x00", name


def test_page_size_and_artwork_match_the_template():
    deck = _deck()
    assert (deck.slide_width, deck.slide_height) == (Pt(1440), Pt(810))      # 1920 x 1080 template
    cover, about, feedback, terms, limits, demands, contact = deck.slides
    for slide in (cover, about, demands, contact):                             # full-page artwork
        pic = _pictures(slide)[0]
        assert (pic.left, pic.top, pic.width, pic.height) == (0, 0, Pt(1440), Pt(810))
    assert len(_pictures(feedback)) == 2 and len(_pictures(terms)) == 1 and len(_pictures(limits)) == 1
    corner = _pictures(terms)[0]
    assert (corner.left, corner.top) == (Pt(-97), Pt(-21))                      # top-left corner graphic
    mirrored = _pictures(feedback)[1]
    assert mirrored.left == Pt(1231) and mirrored._element.spPr.xfrm.get("flipH") == "1"


def test_template_typography():
    deck = _deck()
    cover_title = next(r for r in _runs(deck.slides[0]) if "Aldgate Timber Ltd" in r.text)
    assert cover_title.font.name == "Antonio" and cover_title.font.size == Pt(63.6)
    about = next(r for r in _runs(deck.slides[1]) if r.text == "About Us")
    assert about.font.name == "Poppins" and about.font.bold and about.font.size == Pt(44)
    phone = next(r for r in _runs(deck.slides[1]) if r.text == "0845 3222 525")
    assert phone.font.bold
    contact = [r for r in _runs(deck.slides[6]) if r.text in ("0845 3222 525", "www.ukcreditinsurance.com", "hello@ukcreditinsurance.com")]
    assert len(contact) == 3 and all(r.font.color.rgb == RGBColor(0xFF, 0xFF, 0xFF) for r in contact)


def test_terms_table_geometry_fonts_and_borders():
    deck = _deck()
    terms = next(s for s in deck.slides[3].shapes if s.has_table)
    assert (terms.left, terms.top, terms.width) == (Pt(151), Pt(93), Pt(1126))
    table = terms.table
    assert table.columns[0].width == Pt(370)
    assert not table.first_row and not table.horz_banding                     # no theme banding
    label = table.cell(1, 0).text_frame.paragraphs[0].runs[0]
    value = table.cell(1, 1).text_frame.paragraphs[0].runs[0]
    assert label.font.name == "Poppins" and label.font.bold and label.font.size == Pt(14)
    assert value.font.name == "Arial" and value.font.size == Pt(14)
    tc_pr = table.cell(1, 1)._tc.tcPr
    tags = [child.tag.split("}")[1] for child in tc_pr]
    assert tags[:4] == ["lnL", "lnR", "lnT", "lnB"]                            # borders on every side
    assert table.cell(0, 2).fill.fore_color.rgb == RGBColor(*REC_FILL)

    limits = next(s for s in deck.slides[4].shapes if s.has_table)
    assert (limits.left, limits.top, limits.width) == (Pt(171), Pt(104), Pt(1098))
    buyer = limits.table.cell(1, 0).text_frame.paragraphs[0]
    assert buyer.alignment == 1 and buyer.runs[0].font.name == "Arial"         # left-aligned buyer name
    required = limits.table.cell(1, 2).text_frame.paragraphs[0].runs[0]
    assert required.font.bold and required.text == "£250,000"

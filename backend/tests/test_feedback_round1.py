"""Feedback Round 1 (BRD v1.2 change list): insurer list and naming, money
formatting, Nexus wording rules, credit-limit charges, Additional info as
broker text, Coface limits paths, .xlsm uploads, terms-only / limits-only
decks, the limits sort, expiring policy off the limits page, the static
Important-information block."""

import io
import re

import openpyxl
import pymupdf
import pytest
from pptx import Presentation

from app.core.errors import ExportBlockedError
from app.models.presentation import CreditLimitRow, PresentationColumn, PresentationRequest
from app.models.schemas import (
    BuyerCreditLimit,
    ExtractionResponse,
    ProcessingMeta,
    QuoteExtraction,
    ReviewSummary,
    SetField,
    SetFields,
)
from app.services import apply, grid, library, limits
from app.services.config_store import ConfigInvalid, normalise_terminology
from app.services.money import (
    amount_of,
    format_charges,
    format_money,
    render_value,
    sort_limit_rows,
)
from app.services.presentation import build_limits_xlsx, build_pdf, build_pptx
from app.services.wording import apply_rule, apply_wording_rules
from tests.conftest import make_xlsx, sv
from tests.test_documents_api import _wait_doc
from tests.test_presentation import make_request


def _deck_text(pptx_bytes: bytes) -> str:
    deck = Presentation(io.BytesIO(pptx_bytes))
    parts = []
    for slide in deck.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def _pdf_text(pdf_bytes: bytes) -> str:
    return "\n".join(p.get_text() for p in pymupdf.open(stream=pdf_bytes, filetype="pdf"))


def _extraction(insurer: str | None, *, doc_type="insurer_quote", buyers=(), **values) -> ExtractionResponse:
    fields = {
        name: sv(None, None)
        for name in QuoteExtraction.model_fields
        if name not in ("document_type", "buyer_credit_limits")
    }
    fields["insurer"] = sv(insurer, 1) if insurer else sv(None, None)
    for key, value in values.items():
        fields[key] = sv(value, 1)
    data = QuoteExtraction(document_type=doc_type, buyer_credit_limits=list(buyers), **fields)
    debt, matched = library.debt_collection_rule(insurer)
    return ExtractionResponse(
        meta=ProcessingMeta(filename="doc.pdf", page_count=1, extraction_engine="pymupdf", llm_model="stub"),
        review=ReviewSummary(missing_fields=[], uncertain_fields=[]),
        set_fields=SetFields(debt_collection_support=SetField(value=debt, source="insurer_rule", matched_insurer=matched)),
        data=data,
    )


def _persist(state, doc_id, filename, slot, result):
    """What the upload flow does: the card first, then the extraction folded in."""
    state = apply.upsert_file_entry(state, apply.file_entry(doc_id, filename, slot, status="processing", meta="", job_id=None))
    return apply.apply_result(state, document_id=doc_id, filename=filename, slot=slot, result=result)


def _buyer(name, reg, required, offered):
    return BuyerCreditLimit(buyer_name=name, company_number=reg, limit_required=required, limit_offered=offered, page=1)


# ── A. Insurer list and naming ───────────────────────────────────────────

def test_aviva_is_gone_and_allianz_is_the_display_name():
    names = [i["name"] for i in library.get_insurers()]
    assert "Aviva" not in names and library.match_insurer("Aviva") is None
    assert "Allianz" in names and "Allianz Trade" not in names
    # Identification is unchanged: documents branded "Allianz Trade" still classify.
    assert library.match_insurer("Allianz Trade")["id"] == "allianz"
    assert library.match_insurer("Euler Hermes")["id"] == "allianz"
    assert library.debt_collection_rule("Allianz Trade UK plc") == ("Included", "Allianz")


def test_columns_take_the_standing_display_name():
    state, _ = _persist({}, "d1", "allianz.pdf", "quote",
                                  _extraction("Allianz Trade"))
    state, _ = _persist(state, "d2", "coface.pdf", "quote",
                                  _extraction("Compagnie Française d'Assurance pour le Commerce Extérieur"))
    state, _ = _persist(state, "d3", "x.pdf", "quote",
                                  _extraction("Some Unknown Underwriter"))
    names = [c["name"] for c in state["columns"]]
    assert names == ["Allianz", "Coface", "Some Unknown Underwriter"]
    assert [c["matched"] for c in state["columns"]] == ["Allianz", "Coface", None]
    # The document's own wording is kept on the upload card, not the heading.
    card = next(f for f in state["files"] if f["docId"] == "d2")
    assert "Compagnie" in card["meta"] and "(matched: Coface)" in card["meta"]


def test_cartan_debt_wording_is_config_and_editable():
    assert library.debt_collection_rule("Cartan Trade UK") == ("Inclusive collections", "Cartan")
    assert library.debt_collection_options() == ["Included", "Outsourced", "Inclusive collections"]
    state, _ = _persist({}, "d1", "cartan.pdf", "quote",
                                  _extraction("Cartan Trade"))
    col = state["columns"][0]
    assert col["debt"] == "Inclusive collections"
    cell = grid.cell_view("Whole Turnover", col, "debt")
    assert cell["value"] == "Inclusive collections" and cell["provenance"] == "set"
    # Editable per column to any configured wording, nothing else.
    assert grid.set_override(state, col["id"], "debt", "Outsourced") is not None
    with pytest.raises(grid.InvalidValue):
        grid.set_override(state, col["id"], "debt", "Maybe")


# ── B. Wording, mapping and formatting ───────────────────────────────────

@pytest.mark.parametrize("raw, expected", [
    ("GBP 250,000", "£250,000"),
    ("250000", "£250,000"),
    ("£250,000", "£250,000"),
    ("gbp 55,000.00", "£55,000"),
    ("£1,234.50", "£1,234.50"),
    (" 4,500,000 ", "£4,500,000"),
    ("0", "£0"),                      # an explicit zero is a value
    ("", ""),                         # blank stays blank — never £0
    (None, ""),
    ("Fixed", "Fixed"),               # wording is left verbatim
    ("Nil", "Nil"),
    ("£45 per limit", "£45 per limit"),
    ("0.32%", "0.32%"),
])
def test_money_formatter(raw, expected):
    assert format_money(raw) == expected


@pytest.mark.parametrize("raw, expected", [
    ("£350", "£350"),
    ("GBP 500", "£500"),
    ("£2,300 Fixed Charge - Includes 44 Active Limits and 25 New Requests", "£2,300 / 44 limits"),
    ("£750 Fixed Charge - Includes 30 Credit Limits", "£750 / 30 limits"),
    ("GBP 350 for 25 limits", "£350 / 25 limits"),
    ("£350 / 1 limit", "£350 / 1 limit"),
    ("£2,300 Fixed Charge", "£2,300"),          # no count stated: amount alone
    ("900 for 20 Credit Limits, then £45 each", "£900 / 20 limits"),   # bare leading figure is the amount
    ("20 limits included", "20 limits included"),                    # the figure is the count: verbatim
    ("£45 per limit", "£45 per limit"),         # a rate, not a fixed charge
    ("Included", "Included"),                   # no amount: verbatim
    ("", ""),
])
def test_credit_limit_charges_format(raw, expected):
    assert format_charges(raw) == expected
    assert render_value("credit_limit_charges", raw) == expected


def test_render_value_touches_money_rows_only():
    assert render_value("annual_turnover", "GBP 16,000,000") == "£16,000,000"
    assert render_value("premium_rate", "0.32%") == "0.32%"
    assert render_value("excess_type", "Minimum Retention") == "Minimum Retention"
    assert amount_of("GBP 12,600") == 12600 and amount_of("Nil") == 0 and amount_of("") is None


NEXUS_RULES = library.wording_rules()["nexus"]["max_terms_of_payment"]


@pytest.mark.parametrize("sentence, expected", [
    ("Open credit not exceeding 60 days from end of month of invoice", "60 days end of month"),
    ("Open Credit not exceeding 30 days from the end of the month", "30 days end of month"),
    ("open credit terms of 90 days from end of month", "90 days end of month"),
    ("Open credit not exceeding 45 days from date of invoice", "45 days date of invoice"),
    ("Open credit not exceeding 120 days from invoice date", "120 days date of invoice"),
    ("60 days", "60 days"),                                         # nothing to normalise
    ("Cash against documents", "Cash against documents"),
])
def test_nexus_open_credit_wording_rules(sentence, expected):
    assert apply_rule(sentence, NEXUS_RULES) == expected


def test_wording_rules_apply_per_insurer_after_verification_and_keep_the_page():
    ex = _extraction("Nexus Trade Credit", max_terms_of_payment="Open credit not exceeding 60 days from end of month").data
    changed = apply_wording_rules(ex, "nexus")
    assert changed == ["max_terms_of_payment"]
    assert ex.max_terms_of_payment.value == "60 days end of month"
    assert ex.max_terms_of_payment.page == 1 and ex.max_terms_of_payment.confidence == "high"
    other = _extraction("Atradius", max_terms_of_payment="Open credit not exceeding 60 days from end of month").data
    assert apply_wording_rules(other, "atradius") == []          # another insurer's sentence stays as written


def test_wording_rules_are_validated_configuration():
    ids = {i["id"] for i in library.get_insurers()}
    ok = normalise_terminology({"fields": {}, "rules": {"nexus": {"max_terms_of_payment": [
        {"match": r"(?P<days>\d+) days", "render": "{days} days end of month"}]}}}, ids)
    assert ok["rules"]["nexus"]["max_terms_of_payment"][0]["render"] == "{days} days end of month"
    with pytest.raises(ConfigInvalid):
        normalise_terminology({"fields": {}, "rules": {"nexus": {"max_terms_of_payment": [
            {"match": r"(?P<days>\d+) days", "render": "{weeks} weeks"}]}}}, ids)
    with pytest.raises(ConfigInvalid):
        normalise_terminology({"fields": {}, "rules": {"nexus": {"max_terms_of_payment": [
            {"match": r"(?P<days>\d+ days", "render": "{days}"}]}}}, ids)
    with pytest.raises(ConfigInvalid):
        normalise_terminology({"fields": {}, "rules": {"aviva": {"max_terms_of_payment": []}}}, ids)


def test_additional_info_is_broker_text_not_extraction():
    assert "additional_info" not in QuoteExtraction.model_fields
    assert grid.field("additional_info")["kind"] == "manual"
    from app.llm.stub_extractor import extract_stub
    ex = extract_stub("=== PAGE 1 ===\nINSURER: Atradius\nAdditional Info: countries covered EU\nIndemnity: 90%")
    assert ex.indemnity.value == "90%" and not hasattr(ex, "additional_info")
    # A re-run never touches what the broker typed in the Notes row.
    state, col_id = _persist({}, "d1", "a.pdf", "quote",
                                  _extraction("Atradius", indemnity="90%"))
    state["columns"][0]["data"]["additional_info"] = {"value": "No-claims bonus 10%", "orig": ""}
    state, _ = _persist(state, "d1", "a.pdf", "quote",
                                  _extraction("Atradius", indemnity="85%"))
    assert state["columns"][0]["data"]["additional_info"]["value"] == "No-claims bonus 10%"
    assert state["columns"][0]["data"]["indemnity"]["value"] == "85%"


# ── C. Extraction paths ──────────────────────────────────────────────────

def test_coface_limits_attach_from_a_separate_schedule_and_from_a_quote_addendum():
    # Quote first (the addendum path): buyer rows come with the quote.
    state, col_id = _persist({}, "q1", "coface-quote.pdf", "quote",
                                  _extraction("Coface", buyers=[_buyer("Meridian Foods Ltd", "04821990", "GBP 250,000", "GBP 200,000")]))
    row = state["credit"][0]
    assert row["buyer"] == "Meridian Foods Ltd" and row["offers"][col_id] == "GBP 200,000"

    # A separate schedule whose text names Coface.
    state, _ = _persist(state, "s1", "schedule.xlsx", "limits",
                                  _extraction("Coface", doc_type="credit_limit_schedule",
                                                     buyers=[_buyer("Harbord Retail", "07733120", "120,000", "120,000")]))
    harbord = next(r for r in state["credit"] if r["buyer"] == "Harbord Retail")
    assert harbord["offers"][col_id] == "120,000"

    # A bare buyer list with no insurer in its text: the file name names Coface.
    state, _ = _persist(state, "s2", "Coface limits Aug26.xlsx", "limits",
                                  _extraction(None, doc_type="credit_limit_schedule",
                                                     buyers=[_buyer("Castle Logistics", "09912004", "80,000", "75,000")]))
    castle = next(r for r in state["credit"] if r["buyer"] == "Castle Logistics")
    assert castle["offers"][col_id] == "75,000" and "pending" not in castle
    card = next(f for f in state["files"] if f["docId"] == "s2")
    assert "Coface (from the file name)" in card["meta"]

    # No insurer anywhere: nothing is guessed — the offer waits and the card says so.
    state, _ = _persist(state, "s3", "limits.xlsx", "limits",
                                  _extraction(None, doc_type="credit_limit_schedule",
                                                     buyers=[_buyer("Unknown Buyer Ltd", "11111111", "50,000", "50,000")]))
    unknown = next(r for r in state["credit"] if r["buyer"] == "Unknown Buyer Ltd")
    assert unknown["offers"] == {} and unknown.get("pending", {}) == {}
    assert "insurer not identified" in next(f for f in state["files"] if f["docId"] == "s3")["meta"]


def test_insurer_from_filename_matches_whole_names_only():
    assert apply.insurer_from_filename("Coface limits Aug26.xlsx")["id"] == "coface"
    assert apply.insurer_from_filename("allianz_trade_limits.xlsm")["id"] == "allianz"
    assert apply.insurer_from_filename("straight-limits.xlsx") is None          # "aig" inside a word
    assert apply.insurer_from_filename("limits.xlsx") is None


def test_allianz_xlsm_uploads_and_extracts_limits(client, digital_pdf, monkeypatch):
    from app.services import pipeline as pipeline_mod

    schedule = _extraction("Allianz Trade", doc_type="credit_limit_schedule",
                           buyers=[_buyer("Example Ltd", "01234567", "250000", "200000")]).data
    monkeypatch.setattr(pipeline_mod, "extract_quote_fields", lambda text: schedule)
    assert client.put("/projects/xlsm-p1", json={"client_name": "Xlsm Ltd"}).status_code == 201
    quote = _extraction("Allianz Trade", indemnity="90%").data
    monkeypatch.setattr(pipeline_mod, "extract_quote_fields", lambda text: quote)
    res = client.post("/projects/xlsm-p1/documents",
                      files=[("files", ("q.pdf", digital_pdf, "application/pdf"))], data={"slot": "quote"})
    assert res.status_code == 202
    _wait_doc(client, "xlsm-p1", res.json()["documents"][0]["id"])

    monkeypatch.setattr(pipeline_mod, "extract_quote_fields", lambda text: schedule)
    workbook = make_xlsx({"Limits": [["Buyer", "Reg", "Required", "Agreed"], ["Example Ltd", "01234567", 250000, 200000]]})
    res = client.post(
        "/projects/xlsm-p1/documents",
        files=[("files", ("Allianz limits.xlsm", workbook, "application/vnd.ms-excel.sheet.macroEnabled.12"))],
        data={"slot": "limits"},
    )
    assert res.status_code == 202, res.text
    rec = res.json()["documents"][0]
    assert rec["status"] in ("uploaded", "processing") and rec["content_type"].endswith("macroEnabled.12")
    done = _wait_doc(client, "xlsm-p1", rec["id"])
    assert done["status"] == "ready", done
    state = client.get("/projects/xlsm-p1").json()["state"]
    assert state["credit"][0]["buyer"] == "Example Ltd" and list(state["credit"][0]["offers"].values()) == ["200000"]
    client.delete("/projects/xlsm-p1")


# ── D. Presentation and export ───────────────────────────────────────────

def test_terms_only_limits_only_and_both_all_generate():
    both = make_request()
    assert len(Presentation(io.BytesIO(build_pptx(both))).slides) == 7

    terms_only = make_request(credit_limits=[])
    assert len(Presentation(io.BytesIO(build_pptx(terms_only))).slides) == 6

    limits_only = make_request(columns=[], recommended_id=None, confirmed_fields=[])
    deck = build_pptx(limits_only)                    # no quotes: nothing to confirm
    text = _deck_text(deck)
    assert len(Presentation(io.BytesIO(deck)).slides) == 7
    assert "Terms Comparison" in text and "Credit Limits" in text and "Meridian Foods Ltd" in text
    assert "£250,000" in text
    assert "Allianz Trade, Atradius, Coface were approached but declined" in text
    pdf = _pdf_text(build_pdf(limits_only))
    assert "Credit Limits" in pdf and "Meridian Foods Ltd" in pdf

    with pytest.raises(ExportBlockedError):
        build_pptx(make_request(confirmed_fields=[]))  # with quotes the gate still applies
    with pytest.raises(ValueError):
        PresentationRequest(**dict(make_request().model_dump(), columns=[], credit_limits=[]))


def test_limits_sort_by_limit_required_desc_everywhere():
    rows = [
        CreditLimitRow(buyer="Small", required="£50,000", offers={"a": "£50,000"}),
        CreditLimitRow(buyer="No limit", required="", offers={"a": "£10,000"}),
        CreditLimitRow(buyer="Big", required="GBP 350,000", offers={"a": "£350,000"}),
        CreditLimitRow(buyer="Mid", required="120000", offers={"a": "£120,000"}),
        CreditLimitRow(buyer="Also none", required="", offers={}),
    ]
    assert [r.buyer for r in sort_limit_rows(rows, lambda r: r.required)] == ["Big", "Mid", "Small", "No limit", "Also none"]

    req = make_request(credit_limits=rows, columns=make_request().columns[:1])
    deck_text = _deck_text(build_pptx(req))
    order = [deck_text.index(n) for n in ("Big", "Mid", "Small", "No limit", "Also none")]
    assert order == sorted(order)
    sheet = openpyxl.load_workbook(io.BytesIO(build_limits_xlsx(req))).active
    assert [row[0].value for row in sheet.iter_rows(min_row=2, max_row=6)] == ["Big", "Mid", "Small", "No limit", "Also none"]
    assert sheet["C2"].value == "£350,000" and sheet["C4"].value == "£50,000"

    state = {"columns": [{"id": "a", "name": "Allianz", "manual": False, "expiring": False, "data": {}}],
             "credit": [{"id": str(i), "buyer": r.buyer, "reg": "", "req": r.required, "offers": dict(r.offers)}
                        for i, r in enumerate(rows)]}
    view = limits.grid_view("p", state)
    assert [r["buyer"]["value"] for r in view["rows"]] == ["Big", "Mid", "Small", "No limit", "Also none"]
    headers, body, total = limits._table(view)
    assert [b[0] for b in body] == ["Big", "Mid", "Small", "No limit", "Also none"]
    assert body[0][2] == "£350,000" and body[1][2] == "£120,000"
    assert total[2] == "£520,000"


def test_expiring_policy_stays_on_terms_page_only():
    base = make_request(project_type="renewal")
    expiring = PresentationColumn(id="x", name="Expiring — Allianz Trade",
                                  values={"indemnity": "85%", "annual_turnover": "GBP 4,000,000"})
    req = make_request(project_type="renewal", columns=[expiring, *base.columns],
                       credit_limits=[CreditLimitRow(buyer="Meridian Foods Ltd", required="£250,000",
                                                     offers={"x": "£999", "a": "£250,000", "b": "£200,000"})])
    deck = Presentation(io.BytesIO(build_pptx(req)))
    terms, limits_slide = deck.slides[3], deck.slides[4]
    terms_table = next(s.table for s in terms.shapes if s.has_table)
    assert [c.text for c in terms_table.rows[0].cells] == ["Insurer", "Expiring — Allianz Trade", "Allianz Trade", "Atradius"]
    limits_table = next(s.table for s in limits_slide.shapes if s.has_table)
    assert [c.text for c in limits_table.rows[0].cells] == ["Top Customers", "Company Reg", "Limit Required (GBP)", "Allianz Trade", "Atradius"]
    assert "£999" not in _deck_text(build_pptx(req))
    sheet = openpyxl.load_workbook(io.BytesIO(build_limits_xlsx(req))).active
    assert [c.value for c in sheet[1]] == ["Top Customers", "Company Reg", "Limit Required (GBP)", "Allianz Trade", "Atradius"]
    assert "Expiring" not in _pdf_text(build_pdf(req)).split("Credit Limits")[1].split("Our understanding")[0]


def test_money_is_rendered_as_pounds_on_the_deck_and_blanks_stay_blank():
    col = PresentationColumn(id="q", name="QBE", values={
        "annual_turnover": "GBP 16,000,000", "estimated_annual_premium_exc_ipt": "GBP 55,000",
        "minimum_annual_premium": "", "credit_limit_charges": "GBP 500",
        "excess": "2500", "max_annual_liability": "GBP 2,000,000", "premium_rate": "0.32%",
    })
    req = make_request(columns=[col], recommended_id=None,
                       credit_limits=[CreditLimitRow(buyer="B", required="GBP 80,000", offers={"q": "75000"})])
    for text in (_deck_text(build_pptx(req)), _pdf_text(build_pdf(req))):
        assert re.search(r"GBP\s*\d", text) is None      # "Limit Required (GBP)" is the template's header
        for figure in ("£16,000,000", "£55,000", "£500", "£2,500", "£2,000,000", "£80,000", "£75,000", "0.32%"):
            assert figure in text
        assert "£0" not in text
    table = next(s.table for s in Presentation(io.BytesIO(build_pptx(req))).slides[3].shapes if s.has_table)
    minimum = next(r for r in table.rows if r.cells[0].text.startswith("Minimum Annual Premium"))
    assert minimum.cells[1].text == ""


def test_static_important_information_block_on_every_deck():
    paragraphs = library.presentation_wording()["important_information"]
    assert paragraphs[0].startswith("You will note from the terms comparison")
    assert len(paragraphs) == 5
    for text in (_deck_text(build_pptx(make_request())), _pdf_text(build_pdf(make_request()))):
        for paragraph in paragraphs:
            assert paragraph[:60] in text
        assert "Coface was approached but declined to quote" in text   # the named line stays


def test_presentation_wording_endpoint(client):
    body = client.get("/presentation/wording").json()
    assert "{name}" in body["recommendation"] and len(body["important_information"]) == 5
    assert client.get("/insurers").json()["insurers"][-1] == {
        "id": "cartan", "name": "Cartan", "debt_collection": "included",
        "debt_collection_label": "Inclusive collections",
    }


# ── Maintenance: saved projects follow a display-name change ─────────────

def test_rename_columns_relabels_identified_headings_only(client):
    from app.maintenance import rename_column, rename_columns

    assert rename_column({"id": "1", "name": "Allianz Trade", "matched": "Allianz Trade", "manual": False})["name"] == "Allianz"
    long_name = "Compagnie Française d'Assurance pour le Commerce Extérieur SA. - UK Branch"
    card = [{"docId": "doc-2", "meta": long_name + " (matched: Coface) · quote · 9 pages"}]
    assert rename_column({"id": "2", "name": long_name, "matched": "Coface", "manual": False, "docId": "doc-2"}, card)["name"] == "Coface"
    assert rename_column({"id": "2", "name": long_name, "matched": "Coface", "manual": False, "docId": "doc-2"}) is None  # no card: unknown wording
    assert rename_column({"id": "3", "name": "Expiring — Allianz Trade", "matched": "Allianz Trade", "expiring": True})["name"] == "Expiring — Allianz"
    assert rename_column({"id": "4", "name": "Allianz (option 2)", "matched": "Allianz Trade", "manual": False}) is None  # broker's heading
    assert rename_column({"id": "5", "name": "Free-format column", "manual": True}) is None
    assert rename_column({"id": "6", "name": "Some Unknown Underwriter", "matched": None, "manual": False}) is None
    assert rename_column({"id": "7", "name": "Allianz", "matched": "Allianz", "manual": False, "debt": "Included"}) is None
    # The rule default follows the config change too (Cartan now "Inclusive collections").
    cartan = rename_column({"id": "8", "name": "Cartan", "matched": "Cartan", "manual": False, "debt": "Outsourced",
                            "data": {"debt": {"value": "Outsourced", "orig": ""}}})
    assert cartan["debt"] == "Inclusive collections" and cartan["data"]["debt"]["value"] == "Outsourced"  # override kept

    state = {"columns": [
        {"id": "a", "name": "Allianz Trade", "matched": "Allianz Trade", "manual": False, "data": {}},
        {"id": "b", "name": "My own label", "matched": "Atradius", "manual": False, "data": {}},
    ]}
    assert client.put("/projects/rename-p1", json={"client_name": "Rename Ltd", "state": state}).status_code == 201
    assert rename_columns(dry_run=True) and client.get("/projects/rename-p1").json()["state"]["columns"][0]["name"] == "Allianz Trade"
    changes = rename_columns()
    assert ("rename-p1", "Allianz Trade", "Allianz") in changes
    names = [c["name"] for c in client.get("/projects/rename-p1").json()["state"]["columns"]]
    assert names == ["Allianz", "My own label"]
    assert rename_columns() == []                    # idempotent
    client.delete("/projects/rename-p1")

"""Credit-limits grid (BRD 2.6, S6): rows and per-insurer limits CRUD,
money normalisation, totals, and the editable Excel / PDF downloads."""

import io

import openpyxl
import pymupdf
import pytest

from app.services import limits

PID = "limits-p1"


@pytest.fixture
def project(client):
    state = {
        "columns": [
            {"id": "c-allianz", "name": "Allianz Trade", "manual": False, "expiring": False, "data": {}, "debt": "Included"},
            {"id": "c-qbe", "name": "QBE", "manual": False, "expiring": False, "data": {}, "debt": "Outsourced"},
            {"id": "c-exp", "name": "Expiring policy", "manual": False, "expiring": True, "data": {}, "debt": ""},
        ],
        "credit": [
            {"id": "r1", "buyer": "Example Ltd", "reg": "01234567", "req": "£250,000",
             "offers": {"c-allianz": "£200,000", "c-qbe": "0"}},
            {"id": "r2", "buyer": "Other plc", "reg": "", "req": "£100,000",
             "offers": {"c-allianz": "£100,000"}, "edited": {"c-allianz": True}},
        ],
    }
    body = {"client_name": "Limits Test Ltd", "insurers_approached": ["allianz", "qbe"], "state": state}
    assert client.put(f"/projects/{PID}", json=body).status_code in (200, 201)
    yield client
    client.delete(f"/projects/{PID}")


def _grid(client) -> dict:
    res = client.get(f"/projects/{PID}/limits")
    assert res.status_code == 200, res.text
    return res.json()


def _row(view: dict, row_id: str) -> dict:
    return next(r for r in view["rows"] if r["id"] == row_id)


# ── Money rules (shared with the screen) ────────────────────────────────────

def test_money_rules():
    assert limits.normalise_amount(" 250000 ") == "£250,000"
    assert limits.normalise_amount("£1,234,567.60") == "£1,234,568"
    assert limits.normalise_amount("nil") == "0" and limits.normalise_amount("Declined") == "0"
    assert limits.normalise_amount("Pending") == "Pending"          # wording survives
    assert limits.parse_amount("Pending") is None and limits.is_declined("0")
    assert limits.total(["£1,000", "Pending", "nil", ""]) == 1000 and limits.total(["", "x"]) is None
    assert limits.export_filename("Acme & Sons Ltd!", "pdf") == "Credit Limits - Acme  Sons Ltd.pdf"


# ── Read ────────────────────────────────────────────────────────────────────

def test_grid_view(project):
    view = _grid(project)
    assert [c["id"] for c in view["columns"]] == ["c-allianz", "c-qbe"]      # expiring policy is not a quote
    assert view["hidden_columns"] == [] and view["has_data"] is True
    r1 = _row(view, "r1")
    assert r1["buyer"] == {"value": "Example Ltd", "provenance": "extracted", "amount": None, "declined": False}
    assert r1["required"]["amount"] == 250000
    assert r1["offers"]["c-qbe"] == {"value": "0", "provenance": "extracted", "amount": 0, "declined": True}
    r2 = _row(view, "r2")
    assert r2["offers"]["c-allianz"]["provenance"] == "edited"
    assert r2["offers"]["c-qbe"]["provenance"] == "blank"                     # BRD 2.5: blank ≠ zero
    assert r2["company_number"]["provenance"] == "blank"
    assert view["totals"] == {"required": 350000, "offers": {"c-allianz": 300000, "c-qbe": 0}}
    assert project.get("/projects/none/limits").status_code == 404


# ── Rows ────────────────────────────────────────────────────────────────────

def test_row_crud(project):
    res = project.post(f"/projects/{PID}/limits/rows",
                       json={"buyer": " New Buyer Ltd ", "company_number": "SC123", "required": "75000",
                             "offers": {"c-qbe": "50,000"}})
    assert res.status_code == 201, res.text
    row_id = res.headers["Location"].rsplit("/", 1)[-1]
    row = _row(res.json(), row_id)
    assert row["manual"] is True and row["buyer"]["value"] == "New Buyer Ltd"
    assert row["buyer"]["provenance"] == "manual"
    assert row["required"]["value"] == "£75,000" and row["offers"]["c-qbe"]["value"] == "£50,000"
    assert res.json()["totals"]["offers"]["c-qbe"] == 50000

    res = project.patch(f"/projects/{PID}/limits/rows/r1", json={"required": "300000", "company_number": ""})
    assert res.status_code == 200
    r1 = _row(res.json(), "r1")
    assert r1["required"] == {"value": "£300,000", "provenance": "edited", "amount": 300000, "declined": False}
    assert r1["company_number"]["provenance"] == "blank"
    assert r1["buyer"]["provenance"] == "extracted"                            # untouched field keeps its origin
    assert res.json()["totals"]["required"] == 475000

    assert project.patch(f"/projects/{PID}/limits/rows/nope", json={"buyer": "x"}).status_code == 404
    assert project.post(f"/projects/{PID}/limits/rows", json={"offers": {"c-exp": "£1"}}).status_code == 404
    assert project.post(f"/projects/{PID}/limits/rows", json={"required": "x" * 81}).status_code == 422

    assert project.delete(f"/projects/{PID}/limits/rows/{row_id}").status_code == 204
    assert project.delete(f"/projects/{PID}/limits/rows/{row_id}").status_code == 404
    assert [r["id"] for r in _grid(project)["rows"]] == ["r1", "r2"]

    # The persisted working document has the same shape the screen writes.
    state = project.get(f"/projects/{PID}").json()["state"]
    stored = next(r for r in state["credit"] if r["id"] == "r1")
    assert stored["req"] == "£300,000" and stored["edited"] == {"req": True, "reg": True}
    entries = project.get("/audit").json()["entries"]
    assert any(e["action"] == "limits.row.edit" and e["target"] == PID for e in entries)


# ── Per-insurer limits ──────────────────────────────────────────────────────

def test_offers(project):
    res = project.put(f"/projects/{PID}/limits/rows/r2/offers/c-qbe", json={"value": "nil"})
    assert res.status_code == 200, res.text
    cell = _row(res.json(), "r2")["offers"]["c-qbe"]
    assert cell == {"value": "0", "provenance": "edited", "amount": 0, "declined": True}

    res = project.put(f"/projects/{PID}/limits/rows/r2/offers/c-qbe", json={"value": "Pending"})
    assert _row(res.json(), "r2")["offers"]["c-qbe"]["value"] == "Pending"
    assert res.json()["totals"]["offers"]["c-qbe"] == 0                        # unparseable text is ignored

    res = project.delete(f"/projects/{PID}/limits/rows/r2/offers/c-qbe")
    assert _row(res.json(), "r2")["offers"]["c-qbe"]["provenance"] == "blank"
    assert "c-qbe" not in next(r for r in project.get(f"/projects/{PID}").json()["state"]["credit"]
                               if r["id"] == "r2")["offers"]

    assert project.put(f"/projects/{PID}/limits/rows/r2/offers/c-exp", json={"value": "£1"}).status_code == 404
    assert project.put(f"/projects/{PID}/limits/rows/r2/offers/c-none", json={"value": "£1"}).status_code == 404
    assert project.put(f"/projects/{PID}/limits/rows/none/offers/c-qbe", json={"value": "£1"}).status_code == 404


def test_hide_and_restore_an_insurer_column(project):
    res = project.put(f"/projects/{PID}/limits/columns/c-qbe", json={"hidden": True})
    assert res.status_code == 200
    view = res.json()
    assert [c["id"] for c in view["columns"]] == ["c-allianz"]
    assert view["hidden_columns"] == [{"id": "c-qbe", "name": "QBE"}]
    assert "c-qbe" not in _row(view, "r1")["offers"]
    assert project.put(f"/projects/{PID}/limits/rows/r1/offers/c-qbe", json={"value": "£1"}).status_code == 404
    state = project.get(f"/projects/{PID}").json()["state"]
    assert state["limitsHidden"] == ["c-qbe"]
    assert "c-qbe" not in next(r for r in state["credit"] if r["id"] == "r1")["offers"]   # offers dropped

    res = project.put(f"/projects/{PID}/limits/columns/c-qbe", json={"hidden": False})
    assert [c["id"] for c in res.json()["columns"]] == ["c-allianz", "c-qbe"]
    assert res.json()["hidden_columns"] == []
    assert project.put(f"/projects/{PID}/limits/columns/c-exp", json={"hidden": True}).status_code == 404


# ── Downloads ───────────────────────────────────────────────────────────────

def test_excel_download_is_editable_with_formula_totals(project):
    res = project.get(f"/projects/{PID}/limits/export/xlsx")
    assert res.status_code == 200
    assert res.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert res.headers["content-disposition"] == 'attachment; filename="Credit Limits - Limits Test Ltd.xlsx"'
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    ws = wb["Credit Limits"]
    assert ws["A1"].value == "Credit limits — Limits Test Ltd"
    assert [c.value for c in ws[2]] == ["Top Customers", "Company Reg", "Limit Required (GBP)", "Allianz Trade", "QBE"]
    assert [c.value for c in ws[3]] == ["Example Ltd", "01234567", "£250,000", "£200,000", "0"]
    assert [c.value for c in ws[4]][:4] == ["Other plc", None, "£100,000", "£100,000"]
    assert ws["A5"].value == "Total" and str(ws["C5"].value).startswith("=SUMPRODUCT")
    assert "C3:C5" not in str(ws["C5"].value) and "C3:C4" in str(ws["C5"].value)


def test_pdf_download_has_form_fields_per_value(project):
    res = project.get(f"/projects/{PID}/limits/export/pdf")
    assert res.status_code == 200 and res.headers["content-type"] == "application/pdf"
    assert res.headers["content-disposition"] == 'attachment; filename="Credit Limits - Limits Test Ltd.pdf"'
    doc = pymupdf.open(stream=res.content, filetype="pdf")
    assert doc.page_count == 1
    page = doc[0]
    text = page.get_text()
    assert "Credit limits" in text and "Allianz Trade" in text and "QBE" in text
    assert "Total" in text and "£350,000" in text and "£300,000" in text     # printed totals
    fields = list(page.widgets())
    assert len(fields) == 2 * 5                                              # 2 buyer rows × 5 columns
    values = [w.field_value for w in fields]
    assert "Example Ltd" in values and "01234567" in values
    assert "250,000" in values and "200,000" in values and "0" in values     # amounts editable, GBP in header
    assert all(w.field_type == pymupdf.PDF_WIDGET_TYPE_TEXT for w in fields)


def test_pdf_paginates_long_tables(project):
    for i in range(45):
        project.post(f"/projects/{PID}/limits/rows", json={"buyer": f"Buyer {i}", "required": "1000"})
    res = project.get(f"/projects/{PID}/limits/export/pdf")
    doc = pymupdf.open(stream=res.content, filetype="pdf")
    assert doc.page_count == 3                                               # 47 rows, 20 per page
    assert "Total" in doc[2].get_text() and "Total" not in doc[0].get_text()
    assert sum(len(list(p.widgets())) for p in doc) == 47 * 5


def test_downloads_refuse_an_empty_table(client):
    assert client.put("/projects/limits-empty", json={"client_name": "Empty", "state": {"credit": [
        {"id": "e1", "buyer": "", "reg": "", "req": "", "offers": {}, "manual": True}]}}).status_code in (200, 201)
    try:
        assert client.get("/projects/limits-empty/limits").json()["has_data"] is False
        assert client.get("/projects/limits-empty/limits/export/xlsx").status_code == 409
        assert client.get("/projects/limits-empty/limits/export/pdf").status_code == 409
        assert client.get("/projects/limits-empty/limits/export/docx").status_code == 422
    finally:
        client.delete("/projects/limits-empty")


def test_limits_endpoints_require_sign_in(anon_client):
    assert anon_client.get(f"/projects/{PID}/limits").status_code == 401
    assert anon_client.post(f"/projects/{PID}/limits/rows", json={}).status_code == 401
    assert anon_client.get(f"/projects/{PID}/limits/export/pdf").status_code == 401

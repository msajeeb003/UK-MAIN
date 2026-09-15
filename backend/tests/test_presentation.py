"""Presentation generation (BRD 2.8) — gate, page set, formats, declined line."""

import io

import openpyxl
import pymupdf
import pytest
from pptx import Presentation

from app.core.errors import ExportBlockedError, InvalidDocumentError
from app.models.presentation import (
    CreditLimitRow,
    PresentationColumn,
    PresentationRequest,
)
from app.models.schemas import CONFIRM_REQUIRED_FIELDS
from app.services.presentation import (
    build_limits_xlsx,
    build_pdf,
    build_pptx,
    declined_insurers,
)

ALL_CONFIRMED = list(CONFIRM_REQUIRED_FIELDS)


def make_request(**overrides) -> PresentationRequest:
    base = dict(
        client_name="Aldgate Timber Ltd",
        reference="UKCIB-2418",
        project_type="new",
        columns=[
            PresentationColumn(id="a", name="Allianz Trade", values={
                "type": "Whole Turnover", "annual_turnover": "£4,500,000",
                "premium_rate": "0.32%", "estimated_annual_premium_exc_ipt": "£14,400",
                "indemnity": "90%", "excess": "£1,000",
                "excess_type": "Minimum Retention", "debt": "Included",
            }),
            PresentationColumn(id="b", name="Atradius", values={
                "type": "Whole Turnover", "annual_turnover": "£4,500,000",
                "premium_rate": "0.28%", "estimated_annual_premium_exc_ipt": "£12,600",
                "indemnity": "90%", "excess": "£2,500",
                "excess_type": "Deductible", "debt": "Included",
            }),
        ],
        recommended_id="b",
        approached_insurers=["Allianz Trade", "Atradius", "Coface"],
        credit_limits=[
            CreditLimitRow(buyer="Meridian Foods Ltd", company_number="04821990",
                           required="£250,000", offers={"a": "£250,000", "b": "£200,000"}),
        ],
        notes="All quotes exclude IPT.",
        reasons="Best combination of rate and cover.",
        confirmed_fields=ALL_CONFIRMED,
    )
    base.update(overrides)
    return PresentationRequest(**base)


# ── Gate (BRD 2.5/2.8) ───────────────────────────────────────────────────

def test_gate_blocks_unconfirmed_export():
    req = make_request(confirmed_fields=["indemnity"])
    with pytest.raises(ExportBlockedError, match="3 remaining"):
        build_pptx(req)
    with pytest.raises(ExportBlockedError):
        build_pdf(req)


# ── Declined line (BRD 2.1/S8) ───────────────────────────────────────────

def test_declined_is_approached_minus_quoted():
    req = make_request()
    assert declined_insurers(req) == ["Coface"]


def test_late_quote_moves_off_the_declined_line():
    req = make_request()
    req.columns.append(PresentationColumn(id="c", name="Coface", values={}))
    assert declined_insurers(req) == []


def test_document_wording_column_is_not_declined():
    """A column titled with the insurer's own entity wording must count as
    that insurer's quote — real cases: 'HCC International Insurance
    Company plc' is Tokio Marine HCC; the Coface UK branch's French name
    does not contain 'Coface' at all."""
    req = make_request(
        approached_insurers=["Tokio Marine HCC", "Coface", "Zurich"],
        columns=[
            PresentationColumn(
                id="h", name="HCC International Insurance Company plc",
                matched="Tokio Marine HCC", values={}),
            PresentationColumn(
                id="c",
                name="Compagnie Française d'Assurance pour le Commerce "
                     "Extérieur SA. - UK Branch",
                matched="Coface", values={}),
        ],
        recommended_id="h",
    )
    assert declined_insurers(req) == ["Zurich"]


def test_reasons_starting_with_numbers_stay_intact():
    from app.services.presentation import _reason_lines
    req = make_request(reasons="50% cheaper than the expiring policy\n2) Strong service record")
    assert _reason_lines(req) == [
        "1. 50% cheaper than the expiring policy",
        "2. Strong service record",
    ]


# ── PPTX ─────────────────────────────────────────────────────────────────

def test_pptx_has_the_brd_page_set():
    deck = Presentation(io.BytesIO(build_pptx(make_request())))
    assert len(deck.slides) == 7  # cover, about, info, terms, limits, rec, contact

    all_text = "\n".join(
        shape.text_frame.text
        for slide in deck.slides for shape in slide.shapes
        if shape.has_text_frame
    )
    assert "Aldgate Timber Ltd - Credit Insurance Presentation" in all_text
    # Regulatory wording is always present (BRD 2.7): Duty of Fair
    # Presentation + FCA status, from the approved template.
    assert "Duty of Fair Presentation" in all_text
    assert "Financial Conduct Authority" in all_text
    # Declined auto-line and the merged recommendation name.
    assert "Coface" in all_text and "declined to quote" in all_text
    assert "provided by Atradius" in all_text
    # A blank value never becomes a placeholder.
    assert "N/A" not in all_text


def test_pptx_renewal_title_and_omitted_limits():
    req = make_request(project_type="renewal", credit_limits=[])
    deck = Presentation(io.BytesIO(build_pptx(req)))
    assert len(deck.slides) == 6  # credit-limit page omitted cleanly
    covers = deck.slides[0]
    text = "\n".join(s.text_frame.text for s in covers.shapes if s.has_text_frame)
    assert "Renewal Credit Insurance Presentation" in text


def test_pptx_recommended_column_is_highlighted_and_unselect_clears_it():
    """BRD 2.7: the recommended insurer's column carries the highlight fill in
    the exported comparison table; with no recommendation, no column does."""
    from pptx.dml.color import RGBColor

    from app.services.presentation import REC_FILL

    rec_rgb = RGBColor(*REC_FILL)

    def terms_header_fills(req) -> list:
        deck = Presentation(io.BytesIO(build_pptx(req)))
        table = next(s for s in deck.slides[3].shapes if s.has_table).table
        # Skip the "Insurer" label cell; one fill per insurer column.
        return [table.cell(0, c).fill.fore_color.rgb for c in range(1, len(table.columns))]

    # 'b' recommended -> exactly that column is tinted.
    fills = terms_header_fills(make_request(recommended_id="b"))
    assert fills == [f for f in fills if f is not None]  # all cells have a solid fill
    assert rec_rgb in fills
    assert fills.count(rec_rgb) == 1

    # No recommendation -> the highlight is absent from every column.
    assert rec_rgb not in terms_header_fills(make_request(recommended_id=None))


def test_pptx_terms_table_holds_single_and_six_columns():
    for n in (1, 6):
        cols = [PresentationColumn(id=f"c{i}", name=f"Insurer {i}", values={})
                for i in range(n)]
        deck = Presentation(io.BytesIO(build_pptx(make_request(
            columns=cols, recommended_id="c0", credit_limits=[]))))
        terms = deck.slides[3]
        table = next(s for s in terms.shapes if s.has_table).table
        assert len(table.columns) == 1 + n
        assert len(table.rows) == 1 + 15  # header + the BRD 2.3 rows


# ── PDF ──────────────────────────────────────────────────────────────────

def test_pdf_mirrors_the_page_set():
    doc = pymupdf.open(stream=build_pdf(make_request()), filetype="pdf")
    assert doc.page_count == 7
    text = "\n".join(page.get_text() for page in doc)
    doc.close()
    assert "Terms Comparison" in text
    assert "provided by Atradius" in text
    assert "Duty of Fair Presentation" in text
    assert "Meridian Foods Ltd" in text


def test_pdf_keeps_long_cell_values():
    """Regression: fixed-height cells silently dropped long values (a real
    'Waiting Period for Protracted Default' note vanished from the PDF).
    Rows must grow (or the font shrink) so every reviewed value renders."""
    long_note = ("Waiting Period for Protracted Default: Specified in the "
                 "policy schedule as 6 months from the due date of the "
                 "oldest unpaid invoice, subject to the terms of the policy.")
    req = make_request()
    for col in req.columns:
        col.values["additional_info"] = long_note
    doc = pymupdf.open(stream=build_pdf(req), filetype="pdf")
    text = "\n".join(page.get_text() for page in doc)
    doc.close()
    assert "Waiting Period for Protracted Default" in text


def test_pdf_keeps_long_notes_and_reasons():
    """Regression: _pdf_text dropped the whole text when a broker-written
    note or reasons block outgrew its fixed box."""
    req = make_request(
        notes=("All quotes are subject to underwriting and policy terms. " * 8
               + "END-OF-NOTES-MARKER"),
        reasons="\n".join(f"Reason line {i} with meaningful detail attached"
                          for i in range(1, 9)) + "\nFINAL-REASON-MARKER",
    )
    doc = pymupdf.open(stream=build_pdf(req), filetype="pdf")
    text = "\n".join(page.get_text() for page in doc)
    doc.close()
    assert "END-OF-NOTES-MARKER" in text
    assert "FINAL-REASON-MARKER" in text


def test_pdf_omits_limits_page_when_none():
    doc = pymupdf.open(
        stream=build_pdf(make_request(credit_limits=[])), filetype="pdf"
    )
    assert doc.page_count == 6
    doc.close()


# ── Credit-limits Excel (BRD 2.6 output) ─────────────────────────────────

def test_limits_xlsx_round_trips():
    workbook = openpyxl.load_workbook(io.BytesIO(build_limits_xlsx(make_request())))
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == ("Top Customers", "Company Reg", "Limit Required (GBP)",
                       "Allianz Trade", "Atradius")
    assert rows[1] == ("Meridian Foods Ltd", "04821990", "£250,000", "£250,000", "£200,000")
    # The sample deck's Total row, computed from the parseable amounts
    # (openpyxl reads the empty company-reg cell back as None).
    assert rows[2] == ("Total", None, "£250,000", "£250,000", "£200,000")


def test_limits_xlsx_refused_when_empty():
    with pytest.raises(InvalidDocumentError):
        build_limits_xlsx(make_request(credit_limits=[]))


# ── API endpoint ─────────────────────────────────────────────────────────

def test_endpoint_gate_and_download_headers(client):
    payload = make_request().model_dump()

    blocked = dict(payload, confirmed_fields=[])
    res = client.post("/generate-presentation?format=pptx", json=blocked)
    assert res.status_code == 409

    res = client.post("/generate-presentation?format=pptx", json=payload)
    assert res.status_code == 200
    assert res.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.presentationml"
    )
    assert 'filename="Credit Insurance Presentation of Terms - Aldgate Timber Ltd.pptx"' in (
        res.headers["content-disposition"]
    )
    assert Presentation(io.BytesIO(res.content))  # valid, editable pptx

    res = client.post("/generate-presentation?format=pdf", json=payload)
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/pdf"


def _deck_text(deck) -> str:
    return "\n".join(
        sh.text_frame.text for sl in deck.slides for sh in sl.shapes if sh.has_text_frame
    )


def test_no_recommendation_omits_the_paragraph_but_keeps_fixed_wording():
    """S7: continuing without a recommendation must never merge a placeholder
    name into the client deck; the FCA wording still renders."""
    deck = Presentation(io.BytesIO(build_pptx(make_request(recommended_id=None, reasons="x\ny"))))
    text = _deck_text(deck)
    assert "[no insurer selected]" not in text
    assert "Our recommendation:" not in text
    assert "the following points were important" not in text
    assert "Our understanding of your demands & needs" in text
    assert "Our status:" in text

    with_rec = _deck_text(Presentation(io.BytesIO(build_pptx(make_request(reasons="- first\n- second")))))
    assert "provided by Atradius" in with_rec
    assert "1. first" in with_rec and "2. second" in with_rec

    pdf = pymupdf.open(stream=build_pdf(make_request(recommended_id=None)), filetype="pdf")
    pdf_text = "\n".join(pg.get_text() for pg in pdf)
    assert "[no insurer selected]" not in pdf_text
    assert "Our status:" in pdf_text
